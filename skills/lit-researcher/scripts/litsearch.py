#!/usr/bin/env python3
"""Literature research: deterministic stages around a few LLM decisions.

  check     DIR            before searching: PubMed hit counts for the draft concepts + rubric total and verdict
  search    DIR            run the approved protocol on PubMed, Europe PMC and ClinicalTrials.gov (logged)
  chase     DIR --seeds    add references, citing papers (OpenAlex > Semantic Scholar > Europe PMC) and PubMed similar articles
  rank      DIR            rank all records against the question; write the screening list
  retrieve  DIR            fetch text of included records; shortlist sentences per sub-question
  verify    DIR            check the LLM's evidence picks: sentence exists verbatim, every number in it
  facts     DIR            write facts.md: one sentence per verified fact + registry-built references, for refcheck.py
  paragraph DIR            exact_extracts_paragraph.txt: the verified source sentences as one paragraph, checked by code
  summary   DIR            facts_summary.md: each fact + citation + LLM verdict + exact source sentence, for the chat
  draft     DIR            (optional) turn [E#] markers in draft.md into numbered citations + reference list
  workbook  DIR            write lit_review.xlsx (Evidence, Screening, All records, Search log, Protocol, References),
                           review.html (offline review page) and review_artifact.html (to publish on claude.ai)
  import-edits DIR FILE    write review changes saved in the artifact into lit_review.xlsx

DIR holds protocol.json (written by the LLM, approved by the user) and everything the stages produce.
The LLM writes screening_decisions.json, extractions.json and draft.md; code checks all three.
"""
import argparse
import json
import os
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).parent))
from refcheck import (ANCHOR_THRESHOLD, Http, bm25, clean, collapse, fetch_source,  # noqa: E402
                      key, locate, lookup_doi, sentences, strip_tags)

EUTILS = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils"
EPMC = "https://www.ebi.ac.uk/europepmc/webservices/rest"
REASONS = ["off_topic", "wrong_population", "wrong_design", "no_relevant_outcome", "not_primary_data",
           "duplicate", "not_retrievable", "other"]
FIELDS = ["fact", "design", "n", "population", "comparison", "result"]
NUMBER_FIELDS = ["n", "result"]          # numbers must be in the evidence sentence itself
CONTEXT_FIELDS = ["fact", "population", "comparison"]  # numbers must be somewhere in the paper (e.g. "10-20 montage")


def load(d, name, default=None):
    f = Path(d) / name
    if not f.exists():
        if default is not None:
            return default
        sys.exit(f"Missing {f}")
    return json.loads(f.read_text())


def save(d, name, obj):
    (Path(d) / name).write_text(json.dumps(obj, indent=1, ensure_ascii=False))


def http_for(d):
    return Http(Path(d) / "cache")


def safe(rid):
    return re.sub(r"[^A-Za-z0-9._-]", "_", rid)


# ---------------------------------------------------------------- protocol -> queries

def build_queries(p):
    """Each concept is an OR group of synonyms; concepts are ANDed. Explicit queries in the protocol win."""
    explicit = p.get("queries", {})
    concepts = p.get("concepts", [])
    y0, y1 = (p.get("years") or [None, None]) + [None, None][len(p.get("years") or []):]

    def group(c, fmt_term, fmt_mesh):
        parts = [fmt_term(t) for t in c.get("terms", [])] + [fmt_mesh(m) for m in c.get("mesh", [])]
        return "(" + " OR ".join(parts) + ")"

    pubmed = " AND ".join(group(c, lambda t: f'"{t}"[tiab]', lambda m: f'"{m}"[MeSH Terms]') for c in concepts)
    if y0 or y1:
        pubmed += f' AND ("{y0 or 1900}"[dp] : "{y1 or 3000}"[dp])'
    if p.get("humans_only"):
        pubmed += " NOT (animals[mh] NOT humans[mh])"
    epmc = " AND ".join(group(c, lambda t: f'(TITLE:"{t}" OR ABSTRACT:"{t}")', lambda m: f'MESH_HEADING:"{m}"')
                        for c in concepts)
    if y0 or y1:
        epmc += f" AND PUB_YEAR:[{y0 or 1900} TO {y1 or 3000}]"
    trials = " AND ".join("(" + " OR ".join(c.get("terms", [])) + ")" for c in concepts
                          if c.get("in_trials", True))
    return {"pubmed": explicit.get("pubmed") or pubmed,
            "europepmc": explicit.get("europepmc") or epmc,
            "trials": explicit.get("trials", trials)}


# ---------------------------------------------------------------- record normalisation

def record(rid, **kw):
    r = {"id": rid, "pmid": "", "pmcid": "", "doi": "", "title": "", "authors": [], "first_author": "", "year": None,
         "journal": "", "volume": "", "issue": "", "pages": "", "abstract": "", "pub_types": [], "found_by": []}
    r.update({k: v for k, v in kw.items() if v not in (None, "")})
    r["doi"] = (r["doi"] or "").lower().removeprefix("https://doi.org/")
    if r["authors"] and not r["first_author"]:
        r["first_author"] = r["authors"][0].split()[0] if r["authors"][0] else ""
    return r


def pubmed_records(http, pmids):
    out = []
    for i in range(0, len(pmids), 100):
        status, body = http.get(f"{EUTILS}/efetch.fcgi?db=pubmed&retmode=xml&id=" + ",".join(pmids[i:i + 100]))
        if status != 200:
            continue
        for art in ET.fromstring(body).iter("PubmedArticle"):
            mc = art.find("MedlineCitation")
            a = mc.find("Article")
            j = a.find("Journal")
            authors = []
            for au in a.findall("AuthorList/Author"):
                if au.find("LastName") is not None:
                    authors.append(f"{au.findtext('LastName')} {au.findtext('Initials') or ''}".strip())
                elif au.find("CollectiveName") is not None:
                    authors.append(au.findtext("CollectiveName"))
            abstract = " ".join(
                (f"{t.get('Label')}: " if t.get("Label") else "") + "".join(t.itertext())
                for t in a.findall("Abstract/AbstractText"))
            year = j.findtext("JournalIssue/PubDate/Year") or (j.findtext("JournalIssue/PubDate/MedlineDate") or "")[:4]
            ids = {i.get("IdType"): i.text for i in art.findall("PubmedData/ArticleIdList/ArticleId")}
            out.append(record(
                f"PMID:{mc.findtext('PMID')}", pmid=mc.findtext("PMID"), pmcid=ids.get("pmc", ""), doi=ids.get("doi", ""),
                title=collapse("".join(a.find("ArticleTitle").itertext())), authors=authors,
                first_author=authors[0].rsplit(" ", 1)[0] if authors and " " in authors[0] else (authors or [""])[0],
                year=int(year) if year.isdigit() else None,
                journal=j.findtext("ISOAbbreviation") or j.findtext("Title") or "",
                volume=j.findtext("JournalIssue/Volume") or "", issue=j.findtext("JournalIssue/Issue") or "",
                pages=a.findtext("Pagination/MedlinePgn") or "", abstract=collapse(abstract),
                pub_types=[t.text for t in a.findall("PublicationTypeList/PublicationType")]))
    return out


def epmc_record(r):
    ji = r.get("journalInfo") or {}
    authors = [f"{a.get('lastName', '')} {a.get('initials', '')}".strip() if a.get("lastName") else a.get("collectiveName", "")
               for a in (r.get("authorList") or {}).get("author", [])]
    rid = f"PMID:{r['pmid']}" if r.get("pmid") else f"{r.get('source', 'EPMC')}:{r.get('id')}"
    return record(rid, pmid=r.get("pmid", ""), pmcid=r.get("pmcid", ""), doi=r.get("doi", ""),
                  title=strip_tags(r.get("title", "")), authors=authors,
                  first_author=(authors[0].rsplit(" ", 1)[0] if authors and " " in authors[0] else (authors or [""])[0]),
                  year=int(r["pubYear"]) if str(r.get("pubYear", "")).isdigit() else None,
                  journal=(ji.get("journal") or {}).get("isoabbreviation") or (ji.get("journal") or {}).get("title", ""),
                  volume=ji.get("volume", ""), issue=ji.get("issue", ""), pages=r.get("pageInfo", ""),
                  abstract=strip_tags(r.get("abstractText", "")), pub_types=(r.get("pubTypeList") or {}).get("pubType", []))


def trial_records(http, query, limit):
    fields = "NCTId,BriefTitle,OfficialTitle,LeadSponsorName,StartDate,OverallStatus,BriefSummary,PrimaryOutcomeMeasure"
    _, data = http.json(f"https://clinicaltrials.gov/api/v2/studies?pageSize={min(limit, 100)}&fields={fields}"
                        f"&countTotal=true&query.term=" + quote(query))
    out = []
    for s in (data or {}).get("studies", []):
        p = s.get("protocolSection", {})
        ident, status = p.get("identificationModule", {}), p.get("statusModule", {})
        start = (status.get("startDateStruct") or {}).get("date", "")
        outcomes = "; ".join(o.get("measure", "") for o in (p.get("outcomesModule") or {}).get("primaryOutcomes", []))
        out.append(record(
            f"NCT:{ident.get('nctId')}", title=ident.get("briefTitle", ""),
            authors=[(p.get("sponsorCollaboratorsModule") or {}).get("leadSponsor", {}).get("name", "")],
            year=int(start[:4]) if start[:4].isdigit() else None, journal=f"ClinicalTrials.gov {ident.get('nctId')}",
            abstract=collapse(f"{(p.get('descriptionModule') or {}).get('briefSummary', '')} Primary outcome: {outcomes}. "
                              f"Status: {status.get('overallStatus', '')}; start {start}."),
            pub_types=["Trial registration"]))
    return (data or {}).get("totalCount", len(out)), out


def same_work(a, b):
    """Title matches count as duplicates only when nothing contradicts it: same year, no differing DOI/PMID."""
    if a["doi"] and b["doi"] and a["doi"] != b["doi"]:
        return False
    if a["pmid"] and b["pmid"] and a["pmid"] != b["pmid"]:
        return False
    return not (a["year"] and b["year"] and a["year"] != b["year"])


def merge(records, new, source):
    """Deduplicate by DOI, then PMID, then normalised title (+ same_work). Returns how many were new."""
    index = {}
    for r in records.values():
        for k in (r["doi"], r["pmid"], key(r["title"])[:120]):
            if k:
                index[k] = r["id"]
    added = 0
    for r in new:
        hit = next((index[k] for k in (r["doi"], r["pmid"]) if k and k in index), None)
        tk = key(r["title"])[:120]
        if hit is None and tk in index and same_work(records[index[tk]], r):
            hit = index[tk]
        if hit:
            old = records[hit]
            for f, v in r.items():
                if not old.get(f) and v:
                    old[f] = v
            if source not in old["found_by"]:
                old["found_by"].append(source)
            continue
        r["found_by"] = [source]
        records[r["id"]] = r
        for k in (r["doi"], r["pmid"], key(r["title"])[:120]):
            if k:
                index[k] = r["id"]
        added += 1
    return added


# ---------------------------------------------------------------- search

# ---------------------------------------------------------------- question check (rubric in question_rubric.md)

RUBRIC = [("population", "Population"), ("intervention", "Intervention / exposure / test"), ("comparator", "Comparator"),
          ("outcome", "Outcome"), ("type", "Question type"), ("terms", "Terms"), ("feasibility", "Feasibility (code)"),
          ("timeframe", "Time frame and setting")]


def pubmed_count(http, query):
    _, res = http.json(f"{EUTILS}/esearch.fcgi?db=pubmed&retmode=json&retmax=0&term=" + quote(query))
    return int((res or {}).get("esearchresult", {}).get("count", 0))


def feasibility_score(n):
    return 2 if 10 <= n <= 2000 else 1 if 3 <= n <= 9 or 2001 <= n <= 5000 else 0


def cmd_check(a):
    """Hit counts for the draft concepts, then the rubric total and verdict, as question_check.md."""
    d = Path(a.dir)
    p, qc = load(d, "protocol.json"), load(d, "question_check.json")
    http = http_for(d)
    concepts = p.get("concepts", [])
    counts = {"combined": pubmed_count(http, build_queries({**p, "queries": {}})["pubmed"]), "alone": {}, "without": {}}
    _, res = http.json(f"{EPMC}/search?format=json&pageSize=1&query=" + quote(build_queries({**p, "queries": {}})["europepmc"]))
    counts["europepmc"] = int((res or {}).get("hitCount", 0))
    for c in concepts:
        counts["alone"][c["name"]] = pubmed_count(http, build_queries({**p, "queries": {}, "concepts": [c]})["pubmed"])
        if len(concepts) > 1:
            rest = [x for x in concepts if x is not c]
            counts["without"][c["name"]] = pubmed_count(http, build_queries({**p, "queries": {}, "concepts": rest})["pubmed"])
    y1 = datetime.now().year
    sr_q = build_queries({**p, "queries": {}, "years": [y1 - 5, y1]})["pubmed"] + \
        ' AND (systematic review[pt] OR meta-analysis[pt] OR "systematic review"[ti] OR "meta-analysis"[ti])'
    _, res = http.json(f"{EUTILS}/esearch.fcgi?db=pubmed&retmode=json&retmax=5&sort=pub_date&term=" + quote(sr_q))
    es = (res or {}).get("esearchresult", {})
    counts["reviews"] = {"count": int(es.get("count", 0)), "pmids": es.get("idlist", []), "years": [y1 - 5, y1]}
    reviews = pubmed_records(http, counts["reviews"]["pmids"]) if counts["reviews"]["pmids"] else []
    scores = dict(qc.get("scores", {}))
    missing = [k for k, _ in RUBRIC if k != "feasibility" and k not in scores]
    if missing:
        sys.exit("question_check.json has no score for: " + ", ".join(missing))
    n = counts["combined"]
    scores["feasibility"] = {"score": feasibility_score(n), "reason": f"{n:,} PubMed hits for all concepts combined"}
    for k, v in scores.items():
        if v.get("score") not in (0, 1, 2):
            sys.exit(f"score for {k} must be 0, 1 or 2")
    total = sum(scores[k]["score"] for k, _ in RUBRIC)
    blocking = [label for k, label in RUBRIC if k in ("population", "intervention", "outcome") and scores[k]["score"] == 0]
    verdict = ("Needs refining" if total <= 9 or blocking else "Ready" if total >= 14 else "Workable")
    qc.update({"scores": scores, "total": total, "verdict": verdict, "counts": counts})
    save(d, "question_check.json", qc)

    mark = {0: "🔴 0", 1: "🟡 1", 2: "🟢 2"}
    lines = [f"**Question check: {verdict} ({total}/16)**", "", f"*{qc.get('question', p.get('question', ''))}*", "",
             "| Criterion | Score | Why |", "|---|---|---|"]
    lines += [f"| {label} | {mark[scores[k]['score']]} | {scores[k].get('reason', '')} |" for k, label in RUBRIC]
    lines += ["", f"Question type: {qc.get('question_type', '')}" + (f". Best designs: {qc['best_designs']}" if qc.get("best_designs") else "")]
    lines += ["", f"Hits for the draft concepts: PubMed {n:,}, Europe PMC {counts['europepmc']:,}."]
    if counts["without"]:
        lines.append("Dropping one concept would give: " + "; ".join(
            f"without {k} {v:,}" for k, v in counts["without"].items()) + ".")
    rv = counts["reviews"]
    if rv["count"]:
        lines.append(f"Systematic reviews or meta-analyses since {rv['years'][0]}: {rv['count']} (newest first):")
        lines += [f"- {r['first_author']} et al., {r['journal']} {r['year']}: {r['title']} (PMID {r['pmid']})" for r in reviews]
    else:
        lines.append(f"No systematic review or meta-analysis since {rv['years'][0]} matches these concepts.")
    if blocking:
        lines.append(f"Missing from the question: {', '.join(blocking)}.")
    if verdict != "Ready":
        if qc.get("refined_question"):
            lines += ["", f"**Suggested question** (my additions in [brackets]): {qc['refined_question']}"]
        lines.append("The original question can't be searched well, so the search will use the suggested question."
                     if verdict == "Needs refining" else "You can search the suggested question or the original.")
    (d / "question_check.md").write_text("\n".join(lines) + "\n")
    print("\n".join(lines))


def cmd_search(a):
    d = Path(a.dir)
    p = load(d, "protocol.json")
    if not p.get("approved"):
        sys.exit('protocol.json is not approved. The user must review it and set "approved": true first.')
    http = http_for(d)
    q = build_queries(p)
    limit = int(p.get("max_per_source", 200))
    records, log = load(d, "records.json", {}), load(d, "search_log.json", {"searches": []})
    when = datetime.now().isoformat(timespec="seconds")

    _, res = http.json(f"{EUTILS}/esearch.fcgi?db=pubmed&retmode=json&sort=relevance&retmax={limit}&term=" + quote(q["pubmed"]))
    es = (res or {}).get("esearchresult", {})
    found = pubmed_records(http, es.get("idlist", []))
    log["searches"].append({"source": "PubMed", "query": q["pubmed"], "date": when, "hits": int(es.get("count", 0)),
                            "retrieved": len(found), "new": merge(records, found, "PubMed"),
                            "translated": es.get("querytranslation", "")})

    found, cursor, total = [], "*", 0
    while len(found) < limit:
        _, res = http.json(f"{EPMC}/search?format=json&resultType=core&pageSize=100&cursorMark={quote(cursor)}"
                           f"&query=" + quote(q["europepmc"]))
        total = (res or {}).get("hitCount", total)
        page = (res or {}).get("resultList", {}).get("result", [])
        found += [epmc_record(r) for r in page]
        nxt = (res or {}).get("nextCursorMark")
        if not page or not nxt or nxt == cursor:
            break
        cursor = nxt
    found = found[:limit]
    log["searches"].append({"source": "Europe PMC", "query": q["europepmc"], "date": when, "hits": total,
                            "retrieved": len(found), "new": merge(records, found, "Europe PMC")})

    if q["trials"]:
        total, found = trial_records(http, q["trials"], limit)
        log["searches"].append({"source": "ClinicalTrials.gov", "query": q["trials"], "date": when, "hits": total,
                                "retrieved": len(found), "new": merge(records, found, "ClinicalTrials.gov")})
    save(d, "records.json", records)
    save(d, "search_log.json", log)
    for s in log["searches"][-3:]:
        print(f"{s['source']}: {s['hits']} hits, {s['retrieved']} retrieved, {s['new']} new", file=sys.stderr)
    print(f"{len(records)} unique records", file=sys.stderr)


# ---------------------------------------------------------------- citation chasing

def openalex_record(w):
    inv = w.get("abstract_inverted_index") or {}
    words = sorted((pos, t) for t, ps in inv.items() for pos in ps)
    ids = w.get("ids") or {}
    pmid = (ids.get("pmid") or "").rsplit("/", 1)[-1]
    bib = w.get("biblio") or {}
    authors = [a.get("author", {}).get("display_name", "") for a in w.get("authorships", [])]
    pages = "-".join(x for x in (bib.get("first_page"), bib.get("last_page")) if x)
    return record(f"PMID:{pmid}" if pmid else f"OA:{w.get('id', '').rsplit('/', 1)[-1]}", pmid=pmid,
                  doi=(w.get("doi") or ""), title=w.get("title") or "", authors=authors,
                  first_author=authors[0].split()[-1] if authors and authors[0] else "",
                  year=w.get("publication_year"),
                  journal=((w.get("primary_location") or {}).get("source") or {}).get("display_name", ""),
                  volume=bib.get("volume") or "", issue=bib.get("issue") or "", pages=pages,
                  abstract=" ".join(t for _, t in words), pub_types=[w.get("type", "")])


S2 = "https://api.semanticscholar.org/graph/v1/paper"
S2_FIELDS = "externalIds,title,year,abstract,venue,authors,publicationTypes"


def s2_record(w):
    ext = w.get("externalIds") or {}
    pmid, authors = str(ext.get("PubMed") or ""), [x.get("name", "") for x in w.get("authors") or []]
    return record(f"PMID:{pmid}" if pmid else f"S2:{w.get('paperId')}", pmid=pmid, doi=ext.get("DOI") or "",
                  title=w.get("title") or "", authors=authors,
                  first_author=authors[0].split()[-1] if authors and authors[0] else "", year=w.get("year"),
                  journal=w.get("venue") or "", abstract=w.get("abstract") or "", pub_types=w.get("publicationTypes") or [])


def openalex_auth():
    """Optional OpenAlex credentials from the environment: OPENALEX_API_KEY (free key at openalex.org) and/or OPENALEX_EMAIL."""
    pairs = (("api_key", "OPENALEX_API_KEY"), ("mailto", "OPENALEX_EMAIL"))
    return "".join(f"&{k}={os.environ[v]}" for k, v in pairs if os.environ.get(v))


_OPENALEX_WARNED = []


def links_openalex(http, r, cap):
    select = "id,doi,title,publication_year,authorships,primary_location,ids,biblio,abstract_inverted_index,type,referenced_works"
    ident = f"doi:{r['doi']}" if r["doi"] else f"pmid:{r['pmid']}"
    auth = openalex_auth()
    status, w = http.json(f"https://api.openalex.org/works/{ident}?select={select}{auth}")
    if not w:
        if status == 429 and not _OPENALEX_WARNED:
            _OPENALEX_WARNED.append(1)
            print("OpenAlex: HTTP 429 (rate/credit limit). Set OPENALEX_API_KEY or OPENALEX_EMAIL to use it; "
                  "falling back to the next source.", file=sys.stderr)
        return None
    refs = [x.rsplit("/", 1)[-1] for x in w.get("referenced_works", [])][:cap]
    found = []
    for i in range(0, len(refs), 50):
        _, res = http.json(f"https://api.openalex.org/works?per-page=50&select={select}"
                           f"&filter=openalex:{'|'.join(refs[i:i + 50])}{auth}")
        found += [openalex_record(x) for x in (res or {}).get("results", [])]
    _, res = http.json(f"https://api.openalex.org/works?per-page={cap}&sort=cited_by_count:desc&select={select}"
                       f"&filter=cites:{w['id'].rsplit('/', 1)[-1]}{auth}")
    if res is None:
        if not _OPENALEX_WARNED:
            _OPENALEX_WARNED.append(1)
            print("OpenAlex: citing-papers query failed (often HTTP 429 without OPENALEX_API_KEY); "
                  "falling back to the next source.", file=sys.stderr)
        return None
    citing = [openalex_record(x) for x in res.get("results", [])]
    return found, citing, len(refs) == cap or len(citing) == cap

def links_s2(http, r, cap):
    ident = f"DOI:{r['doi']}" if r["doi"] else f"PMID:{r['pmid']}"
    _, refs = http.json(f"{S2}/{quote(ident)}/references?limit={cap}&fields={S2_FIELDS}")
    _, cites = http.json(f"{S2}/{quote(ident)}/citations?limit={cap}&fields={S2_FIELDS}")
    if refs is None or cites is None:
        return None
    back = [s2_record(x["citedPaper"]) for x in refs.get("data") or [] if (x.get("citedPaper") or {}).get("title")]
    fwd = [s2_record(x["citingPaper"]) for x in cites.get("data") or [] if (x.get("citingPaper") or {}).get("title")]
    return back, fwd, len(refs.get("data") or []) == cap or len(cites.get("data") or []) == cap


def links_epmc(http, r, cap):
    if not r["pmid"]:
        return None
    out = []
    for kind, field in (("references", "referenceList"), ("citations", "citationList")):
        _, res = http.json(f"{EPMC}/MED/{r['pmid']}/{kind}?format=json&pageSize={cap}")
        if res is None:
            return None
        items = ((res.get(field) or {}).get(kind[:-1]) or [])[:cap]
        out.append(pubmed_records(http, [str(x["id"]) for x in items if x.get("source") == "MED" and x.get("id")]))
    return out[0], out[1], max(len(out[0]), len(out[1])) == cap


def links_pubmed(http, r, cap):
    """References and citing papers from PubMed's own links (PMC-indexed only, so fewer; last fallback)."""
    if not r["pmid"]:
        return None
    out = []
    for linkname in ("pubmed_pubmed_refs", "pubmed_pubmed_citedin"):
        status, res = http.json(f"{EUTILS}/elink.fcgi?dbfrom=pubmed&db=pubmed&linkname={linkname}&retmode=json&id={r['pmid']}")
        if res is None:
            return None
        sets = ((res.get("linksets") or [{}])[0].get("linksetdbs") or [])
        ids = [str(x) for x in (sets[0]["links"] if sets else [])][:cap]
        out.append(pubmed_records(http, ids))
    return out[0], out[1], max(len(out[0]), len(out[1])) == cap


def similar_pubmed(http, pmid, n):
    _, res = http.json(f"{EUTILS}/elink.fcgi?dbfrom=pubmed&db=pubmed&linkname=pubmed_pubmed&retmode=json&id={pmid}")
    sets = ((res or {}).get("linksets") or [{}])[0].get("linksetdbs") or []
    ids = [str(x) for x in (sets[0]["links"] if sets else []) if str(x) != pmid][:n]
    return pubmed_records(http, ids)


def cmd_chase(a):
    """References and citing papers of each seed (OpenAlex, else Semantic Scholar, else Europe PMC, else PubMed links),
    plus PubMed's similar articles, which find new papers that nobody has cited yet."""
    d = Path(a.dir)
    http = http_for(d)
    records, log = load(d, "records.json"), load(d, "search_log.json")
    added_total, chased = 0, []
    sources = [("OpenAlex", links_openalex), ("Semantic Scholar", links_s2), ("Europe PMC", links_epmc),
               ("PubMed links", links_pubmed)]
    for rid in a.seeds.split(","):
        rid = rid.strip()
        r = records.get(rid)
        if not r or not (r["doi"] or r["pmid"]):
            print(f"skip {rid}: unknown record or no DOI/PMID", file=sys.stderr)
            continue
        via, found, citing, capped, failed = "", [], [], False, []
        for name, fn in sources:
            got = fn(http, r, a.max)
            if got is not None:
                via, (found, citing, capped) = name, got
                break
            failed.append(name)
        similar = similar_pubmed(http, r["pmid"], a.similar) if r["pmid"] and a.similar else []
        n_back, n_fwd = merge(records, found, f"cited by {rid}"), merge(records, citing, f"cites {rid}")
        n_sim = merge(records, similar, f"similar to {rid}")
        added_total += n_back + n_fwd + n_sim
        chased.append({"seed": rid, "via": via or "none reachable", "unreachable": failed, "references": len(found),
                       "citing": len(citing), "similar": len(similar), "new": n_back + n_fwd + n_sim,
                       "date": datetime.now().isoformat(timespec="seconds"), "capped": capped, "max": a.max})
        print(f"{rid}: {len(found)} references, {len(citing)} citing via {via or 'no source'}"
              + (f" ({', '.join(failed)} unreachable)" if failed else "")
              + f", {len(similar)} similar (PubMed); {n_back + n_fwd + n_sim} new", file=sys.stderr)
    log.setdefault("chasing", []).extend(chased)
    save(d, "records.json", records)
    save(d, "search_log.json", log)
    print(f"citation chasing added {added_total} new records ({len(records)} total)", file=sys.stderr)


# ---------------------------------------------------------------- ranking and screening list

def question_text(p):
    terms = " ".join(t for c in p.get("concepts", []) for t in c.get("terms", []))
    return " ".join([p.get("question", "")] + [s["text"] for s in p.get("subquestions", [])] + [terms])


def cmd_rank(a):
    d = Path(a.dir)
    p, records = load(d, "protocol.json"), load(d, "records.json")
    ids = list(records)
    docs = [f"{records[i]['title']} {records[i]['title']} {records[i]['abstract']}" for i in ids]
    scores = bm25(question_text(p), docs)
    ranked = sorted(zip(ids, scores), key=lambda x: -x[1])
    cap = int(p.get("max_screen", a.max))
    screen = [{"rank": n + 1, "id": i, "score": round(s, 2), "title": records[i]["title"],
               "year": records[i]["year"], "journal": records[i]["journal"], "pub_types": records[i]["pub_types"],
               "found_by": records[i]["found_by"], "abstract": records[i]["abstract"][:2500]}
              for n, (i, s) in enumerate(ranked[:cap])]
    save(d, "screening_list.json", screen)
    log = load(d, "search_log.json")
    log["ranked"] = {"records": len(ids), "screened": len(screen), "not_screened_low_rank": max(0, len(ids) - cap)}
    save(d, "search_log.json", log)
    tmpl = [{"id": s["id"], "decision": "", "reason": "", "note": ""} for s in screen]
    save(d, "screening_decisions.template.json", tmpl)
    chased = [i for i in ids if any(f.startswith(("cited by", "cites", "similar to")) for f in records[i]["found_by"])]
    in_screen = sum(1 for s in screen if s["id"] in chased)
    log["ranked"]["chased_in_screening"] = in_screen
    save(d, "search_log.json", log)
    print(f"{len(screen)} of {len(ids)} records ranked for screening -> screening_list.json", file=sys.stderr)
    if chased and in_screen < 0.1 * len(chased):
        print(f"WARNING: only {in_screen} of {len(chased)} citation-chased records reached the screening list. "
              f"Many have no abstract in OpenAlex, so they rank low; raise max_screen or screen them by title.",
              file=sys.stderr)


# ---------------------------------------------------------------- retrieval and evidence shortlist

def included_ids(d):
    decisions = load(d, "screening_decisions.json")
    bad = [x for x in decisions if x.get("decision") not in ("include", "exclude", "maybe")
           or (x.get("decision") == "exclude" and x.get("reason") not in REASONS)]
    if bad:
        sys.exit(f"{len(bad)} screening decisions are incomplete or use an unknown reason, e.g. {bad[0]}. "
                 f"Decisions: include/exclude/maybe; exclusion reasons: {', '.join(REASONS)}")
    return [x["id"] for x in decisions if x["decision"] == "include"]


def cmd_retrieve(a):
    d = Path(a.dir)
    (d / "sources").mkdir(exist_ok=True)
    p, records, http = load(d, "protocol.json"), load(d, "records.json"), http_for(d)
    out = []
    for rid in included_ids(d):
        r = records[rid]
        if r["journal"].startswith("ClinicalTrials.gov"):
            level, origin, sents = "registry", r["journal"], sentences(r["abstract"])
        else:
            pseudo = {"num": safe(rid), "status": "verified", "doi": r["doi"],
                      "work": {"doi": r["doi"], "abstract": r["abstract"]}}
            level, origin, sents = fetch_source(http, pseudo, a.sources)
            if pseudo["work"]["doi"] == "" and r["abstract"]:
                level, origin, sents = "abstract", "search record", sentences(r["abstract"])
            if level == "none" and r["abstract"]:
                level, origin, sents = "abstract", "search record", sentences(r["abstract"])
        flags = []
        if r["doi"]:
            work, _ = lookup_doi(http, r["doi"])
            for u in (work or {}).get("updates", []):
                if u["type"] in ("retraction", "withdrawal", "removal", "expression_of_concern"):
                    flags.append(f"{u['type'].replace('_', ' ').upper()} ({u['doi']})")
        if any(t.lower() in ("letter", "comment", "editorial", "erratum", "published erratum") for t in r["pub_types"]):
            flags.append("LETTER/COMMENT (check the text is its own, not the article it discusses)")
        if any("preprint" in t.lower() for t in r["pub_types"]) or rid.startswith("PPR:"):
            flags.append("PREPRINT (not peer reviewed)")
        r["flags"] = flags
        save(d / "sources", f"{safe(rid)}.json", sents)
        (d / "sources" / f"{safe(rid)}.txt").write_text("\n".join(f"[s{i}] {s}" for i, s in enumerate(sents)))
        cands = {}
        for sq in p.get("subquestions", []):
            scores = bm25(sq["text"], sents) if sents else []
            cands[sq["id"]] = [{"sid": i, "text": sents[i]} for i, sc in sorted(enumerate(scores), key=lambda x: -x[1])[:a.top]
                               if sc > 0]
        out.append({"paper_id": rid, "title": r["title"], "year": r["year"], "level": level, "origin": origin,
                    "flags": flags,
                    "source_file": f"sources/{safe(rid)}.txt", "candidates": cands})
        print(f"  {rid}: {level} ({len(sents)} sentences)", file=sys.stderr)
    save(d, "records.json", records)
    save(d, "evidence_candidates.json", out)
    tmpl = [{"id": "E1", "paper_id": out[0]["paper_id"] if out else "", "subq": "Q1", "anchor": "",
             **{f: "" for f in FIELDS}, "note": ""}]
    save(d, "extractions.template.json", tmpl)
    print(f"{len(out)} included papers retrieved -> evidence_candidates.json", file=sys.stderr)


# ---------------------------------------------------------------- verification of extracted facts

_NUM = re.compile(r"(?<![A-Za-z])\d+(?:[.,]\d+)*")


def numbers(text):
    """Numbers as written, with thousands separators removed: '51,454' -> '51454', '0.85' -> '0.85'."""
    out = set()
    for m in _NUM.findall(clean(text or "")):
        v = m.replace(",", "") if re.fullmatch(r"\d{1,3}(?:,\d{3})+", m) else m.replace(",", ".")
        out.add(v.lstrip("0") or "0" if "." not in v else v)
    return out


def cmd_verify(a):
    d = Path(a.dir)
    p, records = load(d, "protocol.json"), load(d, "records.json")
    included = set(included_ids(d))
    rows, problems = [], []
    subq = {s["id"]: s["text"] for s in p.get("subquestions", [])}
    for e in load(d, "extractions.json"):
        rid = e.get("paper_id", "")
        row = {**{k: e.get(k, "") for k in ["id", "paper_id", "subq", "anchor", *FIELDS, "note"]},
               "sentence": "", "sid": "", "match": "", "status": ""}
        if rid not in included:
            row["status"] = "paper not included"
        elif e.get("subq") not in subq:
            row["status"] = "unknown sub-question"
        else:
            src = load(d / "sources", f"{safe(rid)}.json")
            score, sid, text = locate(e.get("anchor", ""), src)
            row["match"] = score
            if not (e.get("result") or "").strip():
                row["status"] = "no result: say what was measured"
            elif len((e.get("fact") or "").split()) < 8:
                row["status"] = "no fact: write one plain sentence stating the finding"
            elif score < ANCHOR_THRESHOLD:
                row["status"] = "quote not found"
            else:
                row.update(sentence=text, sid=f"s{sid}")
                missing = sorted(set().union(*(numbers(e.get(f, "")) for f in NUMBER_FIELDS)) - numbers(text))
                paper = numbers(" ".join(src) + " " + records[rid]["title"])
                elsewhere = sorted(set().union(*(numbers(e.get(f, "")) for f in CONTEXT_FIELDS)) - paper)
                row["status"] = (f"numbers not in sentence: {', '.join(missing)}" if missing else
                                 f"numbers not in the paper: {', '.join(elsewhere)}" if elsewhere else "verified")
        if row["status"] != "verified":
            problems.append(f"{row['id']} ({rid}): {row['status']}")
        rows.append(row)
    save(d, "evidence.json", rows)
    for pr in problems:
        print("PROBLEM", pr, file=sys.stderr)
    print(f"{sum(r['status'] == 'verified' for r in rows)}/{len(rows)} evidence rows verified", file=sys.stderr)
    return 1 if problems and a.strict else 0


# ---------------------------------------------------------------- draft with registry-built references

def vancouver(r):
    authors = [re.sub(r"\.", "", a) for a in r["authors"] if a]
    names = ", ".join(authors[:6]) + (", et al" if len(authors) > 6 else "")
    if r["journal"].startswith("ClinicalTrials.gov"):
        nct = r["journal"].split()[-1]
        return f"{names}. {r['title']}. ClinicalTrials.gov identifier: {nct}. https://clinicaltrials.gov/study/{nct}"
    loc = f"{r['year'] or ''}"
    if r["volume"]:
        loc += f";{r['volume']}" + (f"({r['issue']})" if r["issue"] else "")
    if r["pages"]:
        loc += f":{r['pages']}"
    title = r["title"].rstrip(".")
    out = f"{names}. {title}. {r['journal']}. {loc}."
    if r["doi"]:
        out += f" doi:{r['doi']}"
    elif r["pmid"]:
        out += f" PMID:{r['pmid']}"
    return out


def cmd_draft(a):
    d = Path(a.dir)
    records = load(d, "records.json")
    ev = {e["id"]: e for e in load(d, "evidence.json")}
    text = (d / "draft.md").read_text()
    order, errors = [], []

    def cite(m):
        nums = []
        for eid in re.split(r"\s*,\s*", m.group(1)):
            e = ev.get(eid)
            if e is None:
                errors.append(f"{eid}: no such evidence row")
                continue
            if e["status"] != "verified":
                errors.append(f"{eid}: evidence not verified ({e['status']})")
                continue
            if e["paper_id"] not in order:
                order.append(e["paper_id"])
            nums.append(order.index(e["paper_id"]) + 1)
        nums = sorted(set(nums))
        return "[" + ", ".join(map(str, nums)) + "]" if nums else "[CITATION REMOVED]"

    body = re.sub(r"\[((?:E\d+)(?:\s*,\s*E\d+)*)\]", cite, text)
    uncited = [s for s in sentences(re.sub(r"(?m)^#.*$", "", body))
               if len(s.split()) > 8 and not re.search(r"\[\d", s) and not s.startswith("|")]
    refs = "\n".join(f"{i}. {vancouver(records[rid])}" for i, rid in enumerate(order, 1))
    (d / "draft_final.md").write_text(body.rstrip() + "\n\n## References\n\n" + refs + "\n")
    save(d, "draft_report.json", {"errors": errors, "uncited_sentences": uncited, "references": order})
    for e in errors:
        print("PROBLEM", e, file=sys.stderr)
    for s in uncited:
        print("UNCITED", s[:120], file=sys.stderr)
    print(f"draft_final.md written with {len(order)} references", file=sys.stderr)
    return 1 if errors and a.strict else 0


# ---------------------------------------------------------------- fact list for the ref-checker

def cmd_facts(a):
    """One sentence per verified fact, numbered references built from registry metadata, for refcheck.py."""
    d = Path(a.dir)
    records = load(d, "records.json")
    rows = [e for e in load(d, "evidence.json") if e["status"] == "verified"]
    order, lines, fmap = [], [], []
    for e in rows:
        if e["paper_id"] not in order:
            order.append(e["paper_id"])
        k = order.index(e["paper_id"]) + 1
        claim = re.sub(r"[.\s]+$", "", collapse(e["fact"]))
        lines.append(f"{e['id']} ({e['subq']}) — {claim} [{k}].")
        fmap.append({"eid": e["id"], "ref": k, "claim": claim})
    refs = "\n".join(f"{i}. {vancouver(records[rid])}" for i, rid in enumerate(order, 1))
    (d / "facts.md").write_text("# Verified facts\n\n" + "\n\n".join(lines) + "\n\n## References\n\n" + refs + "\n")
    save(d, "facts_map.json", {"facts": fmap, "references": order})
    print(f"facts.md: {len(lines)} facts from {len(order)} papers (unverified rows left out)", file=sys.stderr)


# ---------------------------------------------------------------- chat summary

_VERDICT_MARK = {"supports": "✅ Supports", "partial": "🟡 Partial", "contradicts": "❌ Contradicts",
                 "not_found": "⚠️ Not found in source", "cannot_assess": "⚪ Cannot assess",
                 "unverified_quote": "❓ Quote not verified"}


def short_cite(r):
    who = r["first_author"] or (r["authors"][0] if r["authors"] else "")
    etal = " et al." if len(r["authors"]) > 1 else ""
    return f"{who}{etal}, {r['journal']} {r['year'] or ''}".strip()


def cmd_paragraph(a):
    """exact_extracts_paragraph.txt: the verified source sentences, copied by code, as one paragraph with [ref] numbers.
    Each sentence is checked against the paper's saved text and against the written file before it counts."""
    d = Path(a.dir)
    ev = [r for r in load(d, "evidence.json", []) if r["status"] == "verified"]
    ref = {f["eid"]: f["ref"] for f in load(d, "facts_map.json", {"facts": []})["facts"]}
    ev.sort(key=lambda r: (r["subq"], int(re.sub(r"\D", "", r["id"]) or 0)))
    parts, seen, failed = [], set(), []
    for r in ev:
        if r["id"] not in ref:
            failed.append(f"{r['id']}: not in the fact list (run facts first)")
            continue
        src = load(d / "sources", f"{safe(r['paper_id'])}.json", [])
        pairs = {src[i] + " " + src[i + 1] for i in range(len(src) - 1)}
        if r["sentence"] not in src and r["sentence"] not in pairs:
            failed.append(f"{r['id']}: sentence does not match the paper's saved text")
            continue
        if r["sentence"] in seen:
            continue
        seen.add(r["sentence"])
        sent = r["sentence"].strip()
        parts.append(f"{sent}{'' if sent.endswith(('.', '?', '!')) else '.'} [{ref[r['id']]}]")
    out = d / "exact_extracts_paragraph.txt"
    text = " ".join(parts) + "\n"
    out.write_text(text)
    missing = [s for s in seen if s.strip() not in out.read_text()]
    if failed or missing:
        out.unlink()
        for f in failed + [f"not found verbatim in the written file: {m[:60]}…" for m in missing]:
            print("MISMATCH " + f, file=sys.stderr)
        sys.exit("Paragraph not written: do not deliver it; report these rows to the user.")
    print(f"{len(parts)} sentences ({len(ev)} verified facts{'; repeated sentences used once' if len(parts) < len(ev) else ''}); "
          f"all match the paper's saved text and appear verbatim in {out}", file=sys.stderr)


def cmd_summary(a):
    """facts_summary.md: every fact with its citation, the LLM's suggested verdict and the exact source sentence.
    Built entirely by code from evidence.json, facts_map.json and the ref-check report, so it can be pasted as is."""
    d = Path(a.dir)
    p, records = load(d, "protocol.json"), load(d, "records.json")
    fmap = load(d, "facts_map.json")
    cands = {c["paper_id"]: c for c in load(d, "evidence_candidates.json", [])}
    rc_rows = load(d / "refcheck", "report.json", {"checks": []})["checks"] if (d / "refcheck" / "report.json").exists() else []
    ev = {e["id"]: e for e in load(d, "evidence.json")}
    level = {"fulltext": "full text", "abstract": "abstract only", "registry": "trial registration", "none": "no text"}
    icon = {"supports": "✅", "partial": "🟡", "contradicts": "❌", "not_found": "⚠️", "cannot_assess": "⚪",
            "unverified_quote": "❓"}
    if not a.full:
        out = [f"**Verified facts: {p.get('question', '')}**",
               "*Each fact: the LLM-suggested verdict (✅ supports · 🟡 partial · ❌ contradicts · ⚠️ not found), "
               "the **LLM fact** as extracted, then the **exact extract**, the sentence copied from the paper by code. "
               "*", ""]
    else:
        out = [f"## Verified facts: {p.get('question', '')}", "",
               "*Fact* = the extraction as checked; *Source* = the sentence copied from the paper by code; "
               "*LLM verdict* = the reference check's suggestion, for you to confirm.", ""]
    for sq in p.get("subquestions", []):
        facts = [f for f in fmap["facts"] if ev[f["eid"]]["subq"] == sq["id"]]
        if not a.full:
            stem = sq["text"][0].upper() + sq["text"][1:]
            out.append(f"**{sq['id']}. {stem}**" + ("" if facts else "  \n*No verified facts.*"))
            for f in facts:
                e = ev[f["eid"]]
                rc = next((r for r in rc_rows if r["ref"] == f["ref"]
                           and (key(r["claim"]) + " ").startswith(key(f["eid"]) + " ")), None)
                abs_ = " *(abstract only)*" if (cands.get(e["paper_id"]) or {}).get("level") == "abstract" else ""
                out.append(f"- {icon.get((rc or {}).get('verdict'), '·')} **{e['id']}** [{f['ref']}] **LLM fact:** "
                           f"{e['fact']}{abs_}")
                out.append(f"  **Exact extract ({e['sid']}):** “{e['sentence']}”")
            out.append("")
            continue
        out += [f"### {sq['id']}. {sq['text'][0].upper() + sq['text'][1:]}", ""]
        if not facts:
            out += ["*No verified facts.*", ""]
            continue
        for f in facts:
            e = ev[f["eid"]]
            rc = next((r for r in rc_rows if r["ref"] == f["ref"]
                       and (key(r["claim"]) + " ").startswith(key(f["eid"]) + " ")), None)
            verdict = _VERDICT_MARK.get((rc or {}).get("verdict"), "— not ref-checked")
            r = records[e["paper_id"]]
            flags = "; ".join(r.get("flags", []))
            src = level.get((cands.get(e["paper_id"]) or {}).get("level"), "")
            out += [f"**{e['id']}** · {verdict} · [{f['ref']}] {short_cite(r)} · {src}" + (f" · ⚠️ {flags}" if flags else ""),
                    f"- **Fact:** {f['claim']}",
                    f"- **Source ({e['sid']}):** “{e['sentence']}”"]
            if (rc or {}).get("note"):
                out.append(f"- **LLM note:** {rc['note']}")
            out.append("")
    if a.full:
        out += ["### References", ""] + [f"{i}. {vancouver(records[rid])}" for i, rid in enumerate(fmap["references"], 1)]
    else:
        out += ["**References:** " + " · ".join(f"[{i}] {short_cite(records[rid])}"
                                               for i, rid in enumerate(fmap["references"], 1))]
    (d / "facts_summary.md").write_text("\n".join(out) + "\n", encoding="utf-8")
    print(f"facts_summary.md: {len(fmap['facts'])} facts", file=sys.stderr)


# ---------------------------------------------------------------- workbook

ABSTRACT_HEADER = "Abstract (copied by code; evidence sentence in bold)"
VERDICTS = ["Accepted", "Partly correct", "Rejected"]
OLD_VERDICTS = {"Correct": "Accepted", "Incorrect": "Rejected"}


def kept_inputs(path):
    """The user's own columns from an earlier lit_review.xlsx, so rebuilding the workbook never loses them."""
    kept = {"Evidence": {}, "Screening": {}, "All records": {}}
    if not path.exists():
        return kept
    from openpyxl import load_workbook
    wb = load_workbook(path)
    spec = {"Evidence": ("E#", ["Your verdict", "Your comment", "Your edited fact"]),
            "Screening": ("Record", ["Your decision", "Your comment"]),
            "All records": ("Record", ["Your decision", "Your comment"])}
    for name, (key_col, cols) in spec.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        head = {c.value: i for i, c in enumerate(ws[1])}
        if key_col not in head:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            vals = tuple(row[head[c]] if c in head else None for c in cols)
            if row[head[key_col]] is not None and any(v not in (None, "") for v in vals):
                vals = tuple(OLD_VERDICTS.get(v, v) for v in vals)
                kept[name][str(row[head[key_col]])] = vals
    return kept


def abstract_cell(abstract, sentence):
    """The whole abstract, with the evidence sentence in bold when code can find it there."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from rapidfuzz import fuzz
    text = clean(abstract or "")
    if not text:
        return "No abstract in the search record."
    if not sentence:
        return text
    hit = fuzz.partial_ratio_alignment(sentence, text)
    if hit.score < ANCHOR_THRESHOLD:
        return "[The evidence sentence is from the full text, not the abstract.] " + text
    a, b = hit.dest_start, hit.dest_end
    plain, strong = InlineFont(rFont="Arial", sz=10), InlineFont(rFont="Arial", sz=10, b=True)
    parts = [TextBlock(plain, text[:a]), TextBlock(strong, text[a:b]), TextBlock(plain, text[b:])]
    return CellRichText([x for x in parts if x.text])


def write_review_page(d):
    """review.html: lit_review.xlsx embedded in a self-contained page (reader/writer from RowDoc)."""
    import base64
    here = Path(__file__).parent
    page = (here / "review_template.html").read_text()
    lib = "\n".join(re.sub(r"^(import .*;|export )", "", (here / "review_lib" / f).read_text(), flags=re.M)
                    for f in ("zip.js", "xlsx.js"))
    data = base64.b64encode((d / "lit_review.xlsx").read_bytes()).decode()
    topic = load(d, "protocol.json", {}).get("title") or d.name.replace("-", " ").replace("_", " ").title()
    page = (page.replace("/*__RUNTIME__*/", lib).replace("__XLSX_BASE64__", data).replace("__TOPIC__", topic)
            .replace("__REVIEW_ID__", re.sub(r"[^A-Za-z0-9_-]+", "-", d.name)))
    (d / "review.html").write_text(page.replace("'__MODE__'", "'file'").replace("__TITLE__", f"{topic} Review"))
    # The artifact version: one reusable artifact for every topic, so its title stays the same;
    # claude.ai wraps the page in its own document, so drop ours.
    art = page.replace("'__MODE__'", "'artifact'").replace("__TITLE__", "Lit Review")
    art = re.sub(r"^<!doctype html>\s*<html[^>]*>\s*<head>\s*(<meta [^>]*>\s*)*", "", art)
    art = art.replace("</head>\n<body>\n", "").replace("</body>\n</html>\n", "")
    (d / "review_artifact.html").write_text(art)
    print(f"review page -> {d / 'review.html'} (and review_artifact.html to publish as an artifact)", file=sys.stderr)


def cmd_import_edits(a):
    """Write review changes saved in the artifact (ArtifactData list of reviews/<dir name>/edits) into lit_review.xlsx."""
    from openpyxl import load_workbook
    d = Path(a.dir)
    raw = json.loads(Path(a.file).read_text())
    rows = raw.get("documents") or raw.get("docs") or raw.get("data") or [] if isinstance(raw, dict) else raw
    wb = load_workbook(d / "lit_review.xlsx", rich_text=True)
    n = 0
    for x in rows:
        x = x.get("data", x) if isinstance(x, dict) else {}
        if x.get("sheet") in wb.sheetnames and x.get("row") and x.get("col"):
            wb[x["sheet"]][f"{x['col']}{int(x['row'])}"] = x.get("value") or None
            n += 1
    wb.save(d / "lit_review.xlsx")
    print(f"imported {n} review changes into {d / 'lit_review.xlsx'}", file=sys.stderr)
    write_review_page(d)


def cmd_workbook(a):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    d = Path(a.dir)
    p, records, log = load(d, "protocol.json"), load(d, "records.json"), load(d, "search_log.json")
    screen, decisions = load(d, "screening_list.json", []), {x["id"]: x for x in load(d, "screening_decisions.json", [])}
    evidence, cands = load(d, "evidence.json", []), {c["paper_id"]: c for c in load(d, "evidence_candidates.json", [])}
    draft = load(d, "facts_map.json", None) or load(d, "draft_report.json", {"references": []})
    rc_rows = load(d / "refcheck", "report.json", {"checks": []})["checks"] if (d / "refcheck" / "report.json").exists() else []
    rc_refs = {r["num"]: r for r in load(d / "refcheck", "references.json", [])} if rc_rows else {}
    rc_by_eid = {}
    for f in draft.get("facts", []):
        hit = next((r for r in rc_rows if r["ref"] == f["ref"] and (key(r["claim"]) + " ").startswith(key(f["eid"]) + " ")), None)
        if hit:
            rc_by_eid[f["eid"]] = hit
    rc_status = {"verified": "Verified", "found_without_doi": "Verified (no DOI)", "metadata_discrepancy": "Metadata differs",
                 "doi_mismatch": "DOI is another paper", "doi_not_found": "DOI not found", "not_found": "Not found"}
    rc_verdict = {"supports": "Supports", "partial": "Partial", "contradicts": "Contradicts",
                  "not_found": "Not found in source", "cannot_assess": "Cannot assess", "unverified_quote": "Quote not verified"}
    kept = kept_inputs(d / "lit_review.xlsx")
    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    you = PatternFill("solid", fgColor="FFF9DB")
    fills = {"code": "1F4E78", "llm": "5B3F8C", "you": "7F6000"}

    def sheet(ws, headers, groups, widths, rows):
        ws.append(headers)
        for cell, g in zip(ws[1], groups):
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=fills[g])
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font, cell.alignment = base, wrap
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for col, g in enumerate(groups, 1):
            if g == "you":
                for r in range(2, ws.max_row + 1):
                    ws.cell(r, col).fill = you

    def color(ws, col, last, mapping):
        for label, c in mapping.items():
            ws.conditional_formatting.add(f"{col}2:{col}{max(last, 2)}", FormulaRule(
                formula=[f'LEFT(${col}2,{len(label)})="{label}"'], fill=PatternFill("solid", bgColor=c)))

    wb = Workbook()
    ws = wb.active
    ws.title = "Evidence"
    level = {"fulltext": "Full text", "abstract": "Abstract only", "registry": "Trial registration", "none": "Nothing"}
    cited = set(draft.get("references", []))
    subq = {s["id"]: s["text"] for s in p.get("subquestions", [])}
    sheet(ws, ["E#", "Sub-question", "Paper", "Year", "Flags", "Text checked", "Evidence sentence (copied from source by code)",
               ABSTRACT_HEADER,
               "Check", "LLM fact", "Design", "n", "Population", "Comparison", "Result", "LLM note", "In fact list",
               "Ref-check: reference", "Ref-check: suggested verdict", "Ref-check: note", "Your verdict", "Your comment",
               "Your edited fact"],
          ["code", "llm", "llm", "code", "code", "code", "code", "code", "code", "llm", "llm", "llm", "llm", "llm", "llm",
           "llm", "code", "code", "llm", "llm", "you", "you", "you"],
          [6, 30, 40, 6, 18, 12, 60, 70, 18, 50, 16, 8, 24, 20, 36, 30, 8, 18, 16, 30, 16, 26, 40],
          [[e["id"], f"{e['subq']}: {subq.get(e['subq'], '')}",
            f"{records[e['paper_id']]['first_author']} {records[e['paper_id']]['year']}. {records[e['paper_id']]['title']}"
            if e["paper_id"] in records else e["paper_id"],
            (records.get(e["paper_id"]) or {}).get("year"), "; ".join((records.get(e["paper_id"]) or {}).get("flags", [])),
            level.get((cands.get(e["paper_id"]) or {}).get("level"), ""),
            e["sentence"], None, "Verified" if e["status"] == "verified" else e["status"],
            *[e[f] for f in FIELDS], e["note"], "Yes" if e["paper_id"] in cited else "No",
            rc_status.get((rc_refs.get((rc_by_eid.get(e["id"]) or {}).get("ref")) or {}).get("status"), ""),
            rc_verdict.get((rc_by_eid.get(e["id"]) or {}).get("verdict"), ""), (rc_by_eid.get(e["id"]) or {}).get("note", ""),
            *kept["Evidence"].get(str(e["id"]), (None, None, None))]
           for e in evidence])
    for n, e in enumerate(evidence, 2):
        ws.cell(n, 8).value = abstract_cell((records.get(e["paper_id"]) or {}).get("abstract", ""), e["sentence"])
    color(ws, "I", ws.max_row, {"Verified": "E2EFDA", "numbers": "F8CBAD", "quote": "F8CBAD", "paper": "F8CBAD"})
    color(ws, "E", ws.max_row, {"RETRACTION": "F8CBAD", "WITHDRAWAL": "F8CBAD", "EXPRESSION": "FFF2CC",
                                "PREPRINT": "FFF2CC"})
    color(ws, "T", ws.max_row, {"Supports": "E2EFDA", "Partial": "FFF2CC", "Not found": "FFF2CC",
                                "Contradicts": "F8CBAD", "Quote": "F8CBAD", "Cannot": "EDEDED"})
    color(ws, "S", ws.max_row, {"Verified": "E2EFDA", "Metadata": "FFF2CC", "DOI": "F8CBAD", "Not found": "F8CBAD"})
    color(ws, "U", ws.max_row, {"Accepted": "E2EFDA", "Partly": "FFF2CC", "Rejected": "F8CBAD"})
    dv = DataValidation(type="list", formula1=f'"{",".join(VERDICTS)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"U2:U{ws.max_row + 200}")

    ws = wb.create_sheet("Screening")
    sheet(ws, ["Rank", "Record", "Year", "Journal", "Type", "Found by", "Relevance score", "Title",
               "LLM decision", "LLM reason", "LLM note", "Your decision", "Your comment", "Abstract (copied by code)"],
          ["code"] * 8 + ["llm"] * 3 + ["you"] * 2 + ["code"],
          [6, 16, 6, 20, 18, 20, 9, 60, 12, 18, 30, 14, 26, 80],
          [[s["rank"], s["id"], s["year"], s["journal"], "; ".join(s["pub_types"][:2]), "; ".join(s["found_by"]),
            s["score"], s["title"], (decisions.get(s["id"]) or {}).get("decision", ""),
            (decisions.get(s["id"]) or {}).get("reason", ""), (decisions.get(s["id"]) or {}).get("note", ""),
            *kept["Screening"].get(s["id"], (None, None)), (records.get(s["id"]) or {}).get("abstract", "")]
           for s in screen])
    color(ws, "I", ws.max_row, {"include": "E2EFDA", "maybe": "FFF2CC", "exclude": "EDEDED"})
    dv = DataValidation(type="list", formula1='"Include,Exclude,Maybe"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{ws.max_row + 200}")

    ws = wb.create_sheet("All records")
    ids = list(records)
    scores = bm25(question_text(p), [f"{records[i]['title']} {records[i]['title']} {records[i]['abstract']}" for i in ids])
    order = [i for i, _ in sorted(zip(ids, scores), key=lambda x: -x[1])]
    screened = {s["id"] for s in screen}

    def status(rid, rank):
        x = decisions.get(rid) or {}
        if rid not in screened:
            return f"Not screened (rank {rank} of {len(ids)})"
        if x.get("decision") == "exclude":
            return f"Excluded: {x.get('reason', '')}"
        return {"include": "Included", "maybe": "Maybe"}.get(x.get("decision"), "Screened, no decision")

    sheet(ws, ["Rank", "Record", "Year", "First author", "Title", "Journal", "Type", "Found by", "DOI", "Has abstract",
               "Status", "Your decision", "Your comment"],
          ["code"] * 11 + ["you"] * 2, [6, 16, 6, 14, 60, 20, 16, 20, 22, 8, 26, 14, 26],
          [[n, rid, records[rid]["year"], records[rid]["first_author"], records[rid]["title"], records[rid]["journal"],
            "; ".join(records[rid]["pub_types"][:2]), "; ".join(records[rid]["found_by"]), records[rid]["doi"],
            "Yes" if records[rid]["abstract"] else "No", status(rid, n), *kept["All records"].get(rid, (None, None))]
           for n, rid in enumerate(order, 1)])
    color(ws, "K", ws.max_row, {"Included": "E2EFDA", "Maybe": "FFF2CC", "Excluded": "EDEDED", "Not screened": "FCE4D6"})
    dv = DataValidation(type="list", formula1='"Screen it,Include,Exclude"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{ws.max_row + 200}")

    ws = wb.create_sheet("Search log")
    ws.append(["Source", "Date", "Hits", "Retrieved", "New after de-duplication", "Exact query"])
    for s in log.get("searches", []):
        ws.append([s["source"], s["date"], s["hits"], s["retrieved"], s["new"], s["query"]])
    for c in log.get("chasing", []):
        total = c["references"] + c["citing"] + c.get("similar", 0)
        ws.append([f"Citation chasing: {c['seed']}", c.get("date", ""), total, total, c["new"],
                   f"{c.get('via', 'OpenAlex')}: {c['references']} references + {c['citing']} citing papers"
                   + (f" (capped at {c['max']} each)" if c.get("capped") else "")
                   + (f"; PubMed similar articles: {c['similar']}" if c.get("similar") else "")
                   + (f"; unreachable: {', '.join(c['unreachable'])}" if c.get("unreachable") else "")])
    ws.append([])
    ranked = log.get("ranked", {})
    n_inc = sum(1 for x in decisions.values() if x.get("decision") == "include")
    flow = [("Unique records after de-duplication", len(records)),
            ("Screened (top ranked)", ranked.get("screened", "")),
            ("Not screened (low relevance rank)", ranked.get("not_screened_low_rank", "")),
            ("Excluded at screening", sum(1 for x in decisions.values() if x.get("decision") == "exclude")),
            ("Maybe", sum(1 for x in decisions.values() if x.get("decision") == "maybe")),
            ("Included", n_inc),
            ("Included with full text", sum(1 for c in cands.values() if c["level"] == "fulltext")),
            ("Included papers with no evidence row", n_inc - len({e["paper_id"] for e in evidence})),
            ("Evidence rows", len(evidence)),
            ("Evidence rows verified", sum(1 for e in evidence if e["status"] == "verified")),
            ("Papers in the verified fact list", len(draft.get("references", [])))]
    ws.append(["Flow", "Count"])
    for row in flow:
        ws.append(list(row))
    reasons = Counter(x.get("reason") for x in decisions.values() if x.get("decision") == "exclude")
    for r, n in reasons.most_common():
        ws.append([f"  excluded: {r}", n])
    for row in ws.iter_rows():
        for cell in row:
            cell.font, cell.alignment = base, wrap
    for col, w in zip("ABCDEF", (34, 20, 10, 22, 14, 100)):
        ws.column_dimensions[col].width = w
    ws["A1"].font = ws["A" + str(len(log.get("searches", [])) + len(log.get("chasing", [])) + 3)].font = bold

    ws = wb.create_sheet("Protocol")
    q = build_queries(p)
    rows = [("Question", p.get("question", "")), ("Approved", str(p.get("approved", False)))]
    qc = load(d, "question_check.json", {})
    if qc.get("scores"):
        rows += [("Question as asked", qc.get("question", "")),
                 ("Question check", f"{qc.get('verdict', '')} ({qc.get('total', '')}/16), type: {qc.get('question_type', '')}")]
        rows += [(f"  {label}", f"{qc['scores'][k]['score']}/2: {qc['scores'][k].get('reason', '')}")
                 for k, label in RUBRIC if k in qc["scores"]]
    rows += [(f"Sub-question {s['id']}", s["text"]) for s in p.get("subquestions", [])]
    rows += [(f"Concept: {c.get('name', '')}", "; ".join(c.get("terms", [])) +
              (f"  | MeSH: {'; '.join(c.get('mesh', []))}" if c.get("mesh") else "")) for c in p.get("concepts", [])]
    rows += [("Include", "; ".join(p.get("include", []))), ("Exclude", "; ".join(p.get("exclude", []))),
             ("Years", " to ".join(str(y) for y in p.get("years", []) if y)),
             ("PubMed query", q["pubmed"]), ("Europe PMC query", q["europepmc"]), ("Trials query", q["trials"])]
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows():
        row[0].font, row[1].font = bold, base
        row[1].alignment = wrap
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 110

    ws = wb.create_sheet("References")
    sheet(ws, ["#", "Record", "Reference (built from registry metadata)"], ["code"] * 3, [5, 16, 120],
          [[i, rid, vancouver(records[rid])] for i, rid in enumerate(draft.get("references", []), 1)])
    wb.calculation.fullCalcOnLoad = True
    wb.save(d / "lit_review.xlsx")
    print(f"workbook -> {d / 'lit_review.xlsx'}", file=sys.stderr)
    write_review_page(d)
    write_ris(d)


def write_ris(d):
    """references.ris for Zotero / EndNote / Mendeley: one entry per included paper, built only from registry
    metadata (records.json) plus the code-verified source sentences (evidence.json) as notes."""
    d = Path(d)
    recs = load(d, "records.json")
    recs = recs if isinstance(recs, dict) else {r["id"]: r for r in recs}
    include = [x["id"] for x in load(d, "screening_decisions.json", []) if x.get("decision") == "include"]
    ev = {}
    for e in load(d, "evidence.json", []):
        if str(e.get("status", "")).lower() == "verified":
            ev.setdefault(e["paper_id"], []).append(e)
    topic = load(d, "protocol.json", {}).get("title", "")
    out, n = [], 0
    for rid in include:
        r = recs.get(rid)
        if not r:
            continue
        n += 1
        out.append("TY  - JOUR" if r.get("journal") else "TY  - GEN")
        out.append("TI  - " + (r.get("title") or "").rstrip("."))
        for au in r.get("authors") or []:
            last, _, ini = au.rpartition(" ")
            out.append("AU  - " + (f"{last}, {ini}" if last and ini.isupper() else au))
        for tag, key in (("PY", "year"), ("JO", "journal"), ("VL", "volume"), ("IS", "issue"), ("SP", "pages")):
            if r.get(key):
                out.append(f"{tag}  - {r[key]}")
        if r.get("doi"):
            out += ["DO  - " + r["doi"], "UR  - https://doi.org/" + r["doi"]]
        if r.get("pmid"):
            out += ["UR  - https://pubmed.ncbi.nlm.nih.gov/" + r["pmid"] + "/", "AN  - PMID:" + r["pmid"]]
        if r.get("abstract"):
            out.append("AB  - " + " ".join(r["abstract"].split()))
        for e in ev.get(rid, []):
            out.append(f"N1  - {e['id']} ({e['subq']}) verified source sentence: " + " ".join(e["sentence"].split()))
        out += ["KW  - lit-researcher"] + (["KW  - " + topic] if topic else []) + ["ER  - ", ""]
    (d / "references.ris").write_text("\n".join(out), encoding="utf-8")
    print(f"references.ris -> {d / 'references.ris'} ({n} papers; import into Zotero with File > Import)", file=sys.stderr)


def cmd_ris(a):
    write_ris(a.dir)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    imp = sub.add_parser("import-edits")
    imp.add_argument("dir")
    imp.add_argument("file", help="JSON from the review artifact's reviews/<dir name>/edits collection (ArtifactData list)")
    for name in ("check", "search", "rank", "retrieve", "verify", "facts", "summary", "paragraph", "draft", "workbook", "chase", "queries", "ris"):
        s = sub.add_parser(name)
        s.add_argument("dir")
        if name == "chase":
            s.add_argument("--seeds", required=True, help="comma-separated record ids, e.g. PMID:123,PMID:456")
            s.add_argument("--max", type=int, default=40, help="max references and citing papers per seed")
            s.add_argument("--similar", type=int, default=20, help="PubMed similar articles per seed (0 = none)")
        if name == "rank":
            s.add_argument("--max", type=int, default=150, help="records to screen (protocol max_screen wins)")
        if name == "retrieve":
            s.add_argument("--sources", help="folder of PDFs you have, named by record id (e.g. PMID_123.pdf)")
            s.add_argument("--top", type=int, default=5, help="candidate sentences per sub-question")
        if name == "summary":
            s.add_argument("--full", action="store_true", help="long form with the exact source sentence for each fact")
        if name in ("verify", "draft"):
            s.add_argument("--strict", action="store_true")
    a = ap.parse_args()
    if a.cmd == "queries":
        print(json.dumps(build_queries(load(a.dir, "protocol.json")), indent=1))
        return
    sys.exit({"import-edits": cmd_import_edits, "check": cmd_check, "search": cmd_search, "chase": cmd_chase, "rank": cmd_rank, "retrieve": cmd_retrieve,
              "verify": cmd_verify, "facts": cmd_facts, "summary": cmd_summary, "paragraph": cmd_paragraph, "draft": cmd_draft, "workbook": cmd_workbook, "ris": cmd_ris}[a.cmd](a) or 0)


if __name__ == "__main__":
    main()
