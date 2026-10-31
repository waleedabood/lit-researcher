#!/usr/bin/env python3
"""Reference checker: deterministic stages around one LLM judgment step.

  prepare  MANUSCRIPT --out DIR   parse citations, verify every reference, fetch source
                                  text, shortlist candidate passages per claim
  finalize DIR                    check the LLM's anchor quotes against the source text,
                                  extract the full sentences, write the report

Between the two, someone (the ref-checker agent) reads DIR/checks.json and writes
DIR/judgments.json: one {id, verdict, anchor, note} per check. The anchor is a short
verbatim excerpt from the source; finalize rejects any anchor it cannot find.
"""
import argparse
import csv
import hashlib
import html
import json
import re
import sys
import time
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from math import log
from pathlib import Path
from urllib.parse import quote

import requests
from rapidfuzz import fuzz

try:
    import pysbd
    _SEGMENTER = pysbd.Segmenter(language="en", clean=False)
except ImportError:
    _SEGMENTER = None

VERDICTS = {"supports", "partial", "contradicts", "not_found", "cannot_assess"}
NEEDS_QUOTE = {"supports", "partial", "contradicts"}
ANCHOR_MIN_WORDS = 5
ANCHOR_THRESHOLD = 90
TITLE_MATCH = 90
TITLE_MISMATCH = 70


# ---------------------------------------------------------------- text handling

_CHARMAP = str.maketrans({
    "\u2018": "'", "\u2019": "'", "\u201a": "'", "\u201b": "'", "\u2032": "'",
    "\u201c": '"', "\u201d": '"', "\u201e": '"', "\u2033": '"',
    "\u2010": "-", "\u2011": "-", "\u2012": "-", "\u2013": "-", "\u2014": "-", "\u2212": "-",
    "\u00a0": " ", "\u2009": " ", "\u202f": " ", "\u00ad": "",
})


def clean(text):
    """Undo PDF/Word damage: ligatures, smart quotes, dashes, line-break hyphenation."""
    text = unicodedata.normalize("NFKC", text).translate(_CHARMAP)
    text = re.sub(r"(\w)-\n\s*([a-z])", r"\1\2", text)
    return text.replace("\r\n", "\n").replace("\r", "\n")


def collapse(text):
    return re.sub(r"\s+", " ", text).strip()


def key(text):
    """Matching form: accents stripped, lowercase, alphanumerics only."""
    text = unicodedata.normalize("NFKD", clean(text))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


_ABBR = r"(?:et al|e\.g|i\.e|cf|vs|Fig|Figs|Tab|Ref|Refs|Eq|No|Dr|approx|ca|resp|Suppl|Vol|pp|Jr|Sr|St)"


def sentences(text):
    text = collapse(text)
    if not text:
        return []
    if _SEGMENTER:
        return [s.strip() for s in _SEGMENTER.segment(text) if s.strip()]
    protected = re.sub(r"\b(" + _ABBR + r")\.", lambda m: m.group(1) + "\x00", text)
    protected = re.sub(r"(\d)\.(\d)", "\\1\x00\\2", protected)
    protected = re.sub(r"\b([A-Z])\.(?=\s?[A-Z]\.)", "\\1\x00", protected)  # initials
    parts = re.split(r"(?<=[.!?])\s+(?=[\"'(\[]?[A-Z0-9])", protected)
    return [p.replace("\x00", ".").strip() for p in parts if p.strip()]


def strip_tags(text):
    # only real tags (letter after "<"): "p<0.001" and "p < 0.001" must survive
    return collapse(html.unescape(re.sub(r"</?[A-Za-z][^<>]{0,300}>", " ", text or "")))


# ---------------------------------------------------------------- reading documents

def read_document(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".pdf":
        from pypdf import PdfReader
        return clean("\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages))
    return clean(path.read_text(encoding="utf-8", errors="replace"))


def _read_docx(path):
    """Paragraphs separated by blank lines; superscript citation numbers become [n]."""
    import docx
    paragraphs = []
    for para in docx.Document(str(path)).paragraphs:
        out, sup = [], []
        for run in para.runs + [None]:
            if run is not None and run.font.superscript:
                sup.append(run.text)
                continue
            if sup:
                s = "".join(sup).strip()
                out.append(f"[{s}]" if re.fullmatch(r"[\d,\s\-\u2013]+", s) else s)
                sup = []
            if run is not None:
                out.append(run.text)
        paragraphs.append("".join(out))
    return clean("\n\n".join(paragraphs))


_REF_HEAD = re.compile(
    r"^[ \t]*(?:#+[ \t]*)?(?:\d+\.?[ \t]*)?(references|bibliography|literature cited|works cited|reference list)[ \t]*:?[ \t]*$",
    re.I | re.M)
_AFTER_REFS = re.compile(
    r"^[ \t]*(?:#+[ \t]*)?(figure legends?|figures?|tables?|appendix|supplementary\b.*)[ \t]*:?[ \t]*$", re.I | re.M)


def split_references(text):
    heads = list(_REF_HEAD.finditer(text))
    if not heads:
        return text, ""
    body, refs = text[:heads[-1].start()], text[heads[-1].end():]
    tail = _AFTER_REFS.search(refs)
    return body, refs[:tail.start()] if tail else refs


# ---------------------------------------------------------------- reference list

_NUM = re.compile(r"^\s*(?:\[(\d{1,3})\]|(\d{1,3})[.)]|(\d{1,3})\s(?=[A-Z]))\s*")
_APA_START = re.compile(r"^[A-Z][\w'\-]+(?: [a-z]{1,3} [A-Z][\w'\-]+)?, (?:[A-Z]\.|[A-Z][a-z]+)")
_DOI = re.compile(r"10\.\d{4,9}/[^\s\"<>]+", re.I)


def parse_references(block):
    lines = [collapse(l) for l in block.split("\n")]
    entries, numbered = [], False
    for line in lines:
        m = _NUM.match(line)
        if m and int(next(g for g in m.groups() if g)) == len(entries) + 1:
            numbered = True
            entries.append(line[m.end():])
        elif numbered and line and entries:
            entries[-1] += " " + line
    if len(entries) < 2:
        numbered = False
        nonblank = [l for l in lines if l]
        if re.search(r"\n\s*\n", block.strip()):
            entries = [collapse(c) for c in re.split(r"\n\s*\n", block) if c.strip()]
        elif nonblank and sum(bool(re.search(r"\b(?:19|20)\d{2}\b", l)) for l in nonblank) >= 0.7 * len(nonblank):
            entries = nonblank  # pasted list, one reference per line
        else:
            entries = []
            for line in filter(None, lines):
                if not entries or _APA_START.match(line):
                    entries.append(line)
                else:
                    entries[-1] += " " + line
    return [reference_record(i + 1, raw) for i, raw in enumerate(entries)], numbered


def find_doi(raw):
    m = _DOI.search(raw)
    if not m:
        return None
    doi = m.group(0).rstrip(".,;")
    while doi.endswith(")") and doi.count("(") < doi.count(")"):
        doi = doi[:-1].rstrip(".,;")
    return doi.lower()


def reference_record(num, raw):
    pmid = re.search(r"\bPMID:?\s*(\d{5,9})|pubmed\.ncbi\.nlm\.nih\.gov/(\d{5,9})", raw, re.I)
    pmcid = re.search(r"\b(PMC\d{4,9})\b", raw)
    first = (re.match(r"\s*([^\W\d_][\w'\-]*(?:\s+[^\W\d_][\w'\-]*){0,3}?)\s+[A-Z]{1,4}\s*[,.]", raw)
             or re.match(r"\s*([^\W\d_][\w'\-]*)", raw))
    return {
        "num": num,
        "raw": raw,
        "doi": find_doi(raw),
        "pmid": (pmid.group(1) or pmid.group(2)) if pmid else None,
        "pmcid": pmcid.group(1) if pmcid else None,
        "years": sorted({int(y) for y in re.findall(r"\b((?:19|20)\d{2})[a-z]?\b",
                                                    re.sub(r"https?://\S+|10\.\d{4,9}/\S+", " ", raw))}),
        "first_author": first.group(1) if first else "",
    }


# ---------------------------------------------------------------- HTTP with on-disk cache

class Http:
    def __init__(self, cache_dir, email=None):
        self.cache = Path(cache_dir)
        self.cache.mkdir(parents=True, exist_ok=True)
        self.session = requests.Session()
        agent = "refcheck/1.0 (reference checker)"
        self.session.headers["User-Agent"] = f"{agent} mailto:{email}" if email else agent

    def get(self, url):
        """Returns (status, text). 200/404 responses are cached; network errors give status 0."""
        f = self.cache / (hashlib.sha1(url.encode()).hexdigest() + ".json")
        if f.exists():
            hit = json.loads(f.read_text())
            return hit["status"], hit["body"]
        status, body = 0, ""
        for attempt in range(3):
            try:
                r = self.session.get(url, timeout=30)
            except requests.RequestException as e:
                status, body = 0, str(e)
            else:
                status, body = r.status_code, r.text
                if status in (200, 404):
                    f.write_text(json.dumps({"status": status, "body": body}))
                    break
                if status == 429 and "openalex" in url:
                    break  # daily credit limit: retrying will not help; callers fall back to another source
            time.sleep(1 + 2 * attempt)
        return status, body

    def json(self, url):
        status, body = self.get(url)
        if status != 200:
            return status, None
        try:
            return status, json.loads(body)
        except ValueError:
            return status, None


# ---------------------------------------------------------------- verification

def crossref_summary(m):
    title = strip_tags((m.get("title") or [""])[0])
    title = re.sub(r"^(retracted|withdrawn)(\s+article)?\s*:\s*", "", title, flags=re.I)
    authors = m.get("author") or []
    first = (authors[0].get("family") or authors[0].get("name") or "") if authors else ""
    years = set()
    for k in ("published-print", "published-online", "issued"):
        parts = (m.get(k) or {}).get("date-parts") or [[None]]
        if parts and parts[0] and parts[0][0]:
            years.add(parts[0][0])
    return {
        "doi": (m.get("DOI") or "").lower(),
        "title": title,
        "first_author": first,
        "years": sorted(years),
        "journal": strip_tags((m.get("container-title") or [""])[0]),
        "updates": [{"type": u.get("type"), "doi": u.get("DOI"), "label": u.get("label")}
                    for u in m.get("updated-by") or []],
        "abstract": strip_tags(m.get("abstract") or ""),
        "registry": "crossref",
    }


def datacite_summary(d):
    a = d.get("attributes") or {}
    creators = a.get("creators") or []
    first = (creators[0].get("familyName") or creators[0].get("name") or "") if creators else ""
    return {
        "doi": (a.get("doi") or "").lower(),
        "title": strip_tags(((a.get("titles") or [{}])[0]).get("title", "")),
        "first_author": first,
        "years": [a["publicationYear"]] if a.get("publicationYear") else [],
        "journal": a.get("publisher") or "",
        "updates": [],
        "abstract": "",
        "registry": "datacite",
    }


def lookup_doi(http, doi):
    """Returns (summary or None, exists: bool|None). exists=None means the lookup itself failed."""
    status, data = http.json("https://api.crossref.org/works/" + quote(doi, safe="/():;._-"))
    if data:
        return crossref_summary(data["message"]), True
    status_dc, data = http.json("https://api.datacite.org/dois/" + quote(doi, safe="/():;._-"))
    if data and data.get("data"):
        return datacite_summary(data["data"]), True
    status_h, data = http.json("https://doi.org/api/handles/" + quote(doi, safe="/():;._-"))
    if data and data.get("responseCode") == 1:
        return None, True
    if 0 in (status, status_h):
        return None, None
    return None, False


def compare(ref, work):
    rk = key(ref["raw"])
    tk = key(work["title"])
    if not tk:
        title_score = 0
    elif len(tk.split()) < 3:
        title_score = 100 if f" {tk} " in f" {rk} " else 0
    else:
        title_score = round(fuzz.partial_ratio(tk, rk))
    author = key(work["first_author"])
    group = re.match(r"[^.]*\b(collaborators|group|consortium|committee|task force|investigators|society|association)\b",
                     ref["raw"], re.I)
    cited_first = key(ref["first_author"])
    if not author or group:
        author_ok = None
    elif cited_first == author or cited_first in author.split():
        author_ok = True
    else:
        author_ok = False  # missing, or present but not first (author order wrong)
    year_ok = None
    if ref["years"] and work["years"]:
        diff = min(abs(a - b) for a in ref["years"] for b in work["years"])
        year_ok = True if diff == 0 else ("close" if diff == 1 else False)
    return {"title_score": title_score, "author_ok": author_ok, "year_ok": year_ok}


def classify(match):
    if match["title_score"] < TITLE_MISMATCH and match["author_ok"] and match["year_ok"]:
        return "verified_no_title"  # the reference text omits the title; author and year agree
    if match["title_score"] < TITLE_MISMATCH:
        return "doi_mismatch"
    if match["title_score"] < TITLE_MATCH or match["author_ok"] is False or match["year_ok"] is False:
        return "metadata_discrepancy"
    return "verified"


def search_crossref(http, ref):
    query = _DOI.sub("", ref["raw"])
    query = re.sub(r"https?://\S+", "", query)[:300]
    _, data = http.json("https://api.crossref.org/works?rows=5&query.bibliographic=" + quote(query))
    best = None
    for item in (data or {}).get("message", {}).get("items", []):
        work = crossref_summary(item)
        match = compare(ref, work)
        if match["title_score"] >= TITLE_MATCH and (best is None or match["title_score"] > best[1]["title_score"]):
            best = (work, match)
    return best


def doi_from_id(http, query):
    """DOI for a Europe PMC query such as 'EXT_ID:123 AND SRC:MED' or 'PMCID:PMC123'."""
    _, data = http.json("https://www.ebi.ac.uk/europepmc/webservices/rest/search?format=json&resultType=lite&query="
                        + quote(query))
    results = (data or {}).get("resultList", {}).get("result", [])
    if not results:
        return None
    return (results[0].get("doi") or "").lower() or None


def trial_record(http, nct):
    """A ClinicalTrials.gov registration in the same shape as a registry work, plus its text as the source."""
    fields = "NCTId,BriefTitle,OfficialTitle,LeadSponsorName,StartDate,OverallStatus,BriefSummary,PrimaryOutcomeMeasure"
    _, data = http.json(f"https://clinicaltrials.gov/api/v2/studies/{nct}?fields={fields}")
    p = (data or {}).get("protocolSection")
    if not p:
        return None
    ident, status = p.get("identificationModule", {}), p.get("statusModule", {})
    start = (status.get("startDateStruct") or {}).get("date", "")
    summary = (p.get("descriptionModule") or {}).get("briefSummary", "")
    outcomes = [o.get("measure", "") for o in (p.get("outcomesModule") or {}).get("primaryOutcomes", [])]
    text = [ident.get("officialTitle", ""), summary] + [f"Primary outcome: {o}." for o in outcomes] + [
        f"Overall status: {status.get('overallStatus', '')}; start date {start}."]
    return {
        "doi": "",
        "title": ident.get("briefTitle", ""),
        "first_author": "",
        "years": [int(start[:4])] if start[:4].isdigit() else [],
        "journal": "ClinicalTrials.gov " + nct,
        "updates": [],
        "abstract": "",
        "registry": "clinicaltrials.gov",
        "trial_text": sentences(" ".join(t for t in text if t)),
    }


def verify_reference(http, ref):
    notes = []
    nct = re.search(r"\bNCT\d{8}\b", ref["raw"])
    if nct and not ref["doi"] and not ref["pmid"]:
        work = trial_record(http, nct.group(0))
        match = compare(ref, work) if work else None
        status = ("not_found" if not work else
                  "verified" if match["title_score"] >= TITLE_MATCH and match["year_ok"] is not False
                  else "metadata_discrepancy")
        if work and match["year_ok"] is False:
            notes.append(f"registry start year {work['years']}")
        ref.update({"status": status, "match": match, "work": work, "retracted": False, "concern": False,
                    "corrected": False, "notes": notes + ["trial registration checked on ClinicalTrials.gov"]})
        return ref
    doi = ref["doi"]
    for label, value, query in (("PMID", ref["pmid"], f"EXT_ID:{ref['pmid']} AND SRC:MED"),
                                ("PMCID", ref.get("pmcid"), f"PMCID:{ref.get('pmcid')}")):
        if not doi and value:
            doi = doi_from_id(http, query)
            if doi:
                notes.append(f"DOI {doi} found from {label} {value}")
    work, match, status = None, None, None
    if doi:
        work, exists = lookup_doi(http, doi)
        if exists is None:
            status = "lookup_failed"
            notes.append("network error while resolving the DOI; rerun to retry")
        elif not exists:
            status = "doi_not_found"
        elif work is None:
            status = "doi_exists_no_metadata"
        else:
            match = compare(ref, work)
            status = classify(match)
            if status == "verified_no_title":
                status = "verified"
                notes.append("reference text has no title; identity rests on the identifier, first author and year "
                             f"(registry title: {work['title']})")
    if status in (None, "doi_not_found", "doi_mismatch"):
        found = search_crossref(http, ref)
        if found:
            if status is None:
                work, match = found
                status = "found_without_doi" if classify(match) == "verified" else "metadata_discrepancy"
            elif found[0]["doi"] != doi:
                notes.append(f"a work matching this reference has DOI {found[0]['doi']}")
                ref["suggested_doi"] = found[0]["doi"]
                ref["suggested_title"] = found[0]["title"]
        elif status is None:
            status = "not_found"
    if match:
        if match["year_ok"] == "close":
            notes.append(f"year differs by one from the registry ({work['years']}); usually online-first vs print")
        if match["year_ok"] is False:
            notes.append(f"cited year {', '.join(map(str, ref['years']))}; registry year {', '.join(map(str, work['years']))}")
        if match["author_ok"] is False:
            notes.append(f"cited first author is {ref['first_author']}; registry first author is {work['first_author']}")
    updates = (work or {}).get("updates", [])
    ref.update({
        "status": status,
        "match": match,
        "work": work,
        "retracted": any(u["type"] in ("retraction", "withdrawal", "removal") for u in updates),
        "concern": any(u["type"] == "expression_of_concern" for u in updates),
        "corrected": any(u["type"] in ("correction", "erratum") for u in updates),
        "notes": notes,
    })
    return ref


# ---------------------------------------------------------------- source text

def _jats_sentences(xml_text):
    root = ET.fromstring(xml_text)
    for tag in ("ref-list", "back"):
        for parent in root.iter():
            for child in list(parent):
                if child.tag == tag:
                    parent.remove(child)
    out = []
    for p in root.iter("p"):
        out.extend(sentences("".join(p.itertext())))
    return out


def fetch_source(http, ref, local_dir):
    """Best available text of the cited work: local file > PMC full text > abstract."""
    num = ref["num"]
    if local_dir:
        for name in (f"{num}", f"ref{num}", f"ref_{num}"):
            for suffix in (".pdf", ".txt", ".md", ".docx"):
                f = Path(local_dir) / (name + suffix)
                if f.exists():
                    body, _ = split_references(read_document(f))
                    return "fulltext", f"local file {f.name}", sentences(body)
    work = ref.get("work") or {}
    if work.get("trial_text"):
        return "registry", f"ClinicalTrials.gov record ({work['journal'].split()[-1]})", work["trial_text"]
    if ref["status"] == "doi_mismatch" and ref.get("suggested_doi"):
        # the cited identifier is another paper; read the paper the citation describes
        doi, work = ref["suggested_doi"], {}
    elif ref["status"] in ("doi_not_found", "not_found", "lookup_failed"):
        doi = None
    else:
        doi = work.get("doi") or ref.get("doi")
    if not doi:
        return "none", "no verified DOI to look up", []
    _, data = http.json("https://www.ebi.ac.uk/europepmc/webservices/rest/search?format=json&resultType=core&query="
                        + quote(f'DOI:"{doi}"'))
    results = (data or {}).get("resultList", {}).get("result", [])
    hit = results[0] if results else {}
    if hit.get("pmcid"):
        status, body = http.get(f"https://www.ebi.ac.uk/europepmc/webservices/rest/{hit['pmcid']}/fullTextXML")
        if status == 200 and body.strip():
            try:
                sents = _jats_sentences(body)
            except ET.ParseError:
                sents = []
            if len(sents) > 20:
                return "fulltext", f"Europe PMC {hit['pmcid']}", sents
    if hit.get("abstractText"):
        return "abstract", "Europe PMC abstract", sentences(strip_tags(hit["abstractText"]))
    if work.get("abstract"):
        return "abstract", "Crossref abstract", sentences(work["abstract"])
    _, data = http.json("https://api.openalex.org/works/doi:" + quote(doi, safe="/():;._-"))
    inverted = (data or {}).get("abstract_inverted_index")
    if inverted:
        words = sorted((pos, w) for w, positions in inverted.items() for pos in positions)
        return "abstract", "OpenAlex abstract", sentences(" ".join(w for _, w in words))
    return "none", "no open full text or abstract found", []


# ---------------------------------------------------------------- citations in the body

_NUMCITE = re.compile(r"\[(\d{1,3}(?:\s*[-,]\s*\d{1,3})*)\]")
_PARENCITE = re.compile(r"\((\d{1,3}(?:\s*[-,]\s*\d{1,3})*)\)")
_AY_PAREN = re.compile(r"\(([^()]*?\b(?:19|20)\d{2}[a-z]?\b[^()]*)\)")
_AY_ITEM = re.compile(r"([^\W\d_][\w'\-]+)(?:\s+et\s+al\.?|\s+(?:and|&)\s+[^\W\d_][\w'\-]+)?,?\s+((?:19|20)\d{2})[a-z]?")
_AY_NARRATIVE = re.compile(r"([^\W\d_][\w'\-]+)(?:\s+et\s+al\.?|\s+(?:and|&)\s+[^\W\d_][\w'\-]+)?\s+\(((?:19|20)\d{2})[a-z]?\)")


def expand(spec):
    out = []
    for part in re.split(r"\s*,\s*", spec.strip()):
        if "-" in part:
            a, b = (int(x) for x in part.split("-")[:2])
            out.extend(range(a, b + 1) if 0 <= b - a <= 50 else [a, b])
        elif part:
            out.append(int(part))
    return out


_CHIP_LINE = re.compile(r"^[ \t]*(?:PubMed|PMC|DOI|Crossref|Google Scholar|Wikipedia|Sources?|\+\d+)[ \t]*$",
                        re.I | re.M)


def body_sentences(body):
    body = _CHIP_LINE.sub("", body)
    # "...false.[3] Next" -> "...false [3]. Next" so the marker stays with its sentence
    body = re.sub(r"([.!?])\s?((?:\[[\d,\s\-]+\])+)(?=\s+[\"'(\[]?[A-Z]|\s*$)", r" \2\1", body, flags=re.M)
    out = []
    lines = body.count("\n")
    splitter = r"\n\s*\n" if body.count("\n\n") * 5 >= lines else r"\n"
    for block in re.split(splitter, body):
        if re.fullmatch(r"\s*#.*", block) or len(block.split()) < 4:
            continue  # headings
        out.extend(sentences(block))
    return out


def find_citations(sents, refs, numbered, unresolved=None):
    by_author_year = {}
    for r in refs:
        surnames = {key(r["first_author"])}
        if r.get("work") and r["work"].get("first_author"):
            surnames.add(key(r["work"]["first_author"]))
        years = set(r["years"]) | set((r.get("work") or {}).get("years", []))
        for s in list(filter(None, surnames)):
            for name in {s, s.split()[-1], s.split()[0]}:
                for y in years:
                    nums = by_author_year.setdefault((name, y), [])
                    if r["num"] not in nums:
                        nums.append(r["num"])
    by_trial = {m: r["num"] for r in refs for m in re.findall(r"NCT\d{8}", r["raw"])}
    use_parens = numbered and not any(_NUMCITE.search(s) for s in sents)
    found = []
    for i, s in enumerate(sents):
        nums = []
        if numbered:
            for m in _NUMCITE.finditer(s):
                nums += expand(m.group(1))
            if use_parens:
                for m in _PARENCITE.finditer(s):
                    cand = expand(m.group(1))
                    if all(1 <= n <= len(refs) for n in cand):
                        nums += cand
        else:
            items = [m for p in _AY_PAREN.finditer(s) for m in _AY_ITEM.finditer(p.group(1))]
            items += list(_AY_NARRATIVE.finditer(s))
            trials = re.findall(r"NCT\d{8}", s)
            for m in items:
                if trials and m.group(1).lower() in ("start", "april", "january", "february", "march", "may", "june",
                                                     "july", "august", "september", "october", "november", "december"):
                    continue  # a date inside "(NCT…, start April 2026)", not an author
                hit = by_author_year.get((key(m.group(1)), int(m.group(2))))
                if hit is None and unresolved is not None:
                    unresolved.append(collapse(m.group(0)))
                nums += hit or [-1]
            nums += [by_trial[t] for t in trials if t in by_trial]
        if nums:
            found.append((i, list(dict.fromkeys(nums))))
    return found


def strip_markers(s):
    s = _NUMCITE.sub("", s)
    s = _AY_PAREN.sub("", s)
    return re.sub(r"\s+([.,;:])", r"\1", collapse(s))


# ---------------------------------------------------------------- shortlist (BM25)

_STOP = set("""a an the of and or in on to for with by from as at is are was were be been being this that these
those it its we our their than then there which who whom whose into also not no but can may might has have had
such more most less other both each between after before during over under within without patients study""".split())


def tokens(s):
    return [t[:-1] if len(t) > 4 and t.endswith("s") else t for t in key(s).split() if t not in _STOP and len(t) > 1]


def bm25(query, docs, k1=1.5, b=0.75):
    doc_tokens = [tokens(d) for d in docs]
    n = len(doc_tokens)
    avg = (sum(map(len, doc_tokens)) / n) or 1
    df = Counter(t for d in doc_tokens for t in set(d))
    q = set(tokens(query))
    scores = []
    for d in doc_tokens:
        tf = Counter(d)
        score = 0.0
        for t in q & tf.keys():
            idf = log(1 + (n - df[t] + 0.5) / (df[t] + 0.5))
            score += idf * tf[t] * (k1 + 1) / (tf[t] + k1 * (1 - b + b * len(d) / avg))
        scores.append(score)
    return scores


# ---------------------------------------------------------------- prepare

def prepare(args):
    out = Path(args.out)
    (out / "sources").mkdir(parents=True, exist_ok=True)
    text = read_document(args.manuscript)
    body, ref_block = split_references(text)
    if args.refs:
        ref_block = read_document(args.refs)
    if not ref_block.strip():
        sys.exit("No reference list found. Add a 'References' heading or pass --refs FILE.")
    refs, numbered = parse_references(ref_block)
    print(f"{len(refs)} references ({'numbered' if numbered else 'author-year'}); verifying...", file=sys.stderr)

    http = Http(out / "cache", args.email)

    def work(ref):
        verify_reference(http, ref)
        level, origin, sents = fetch_source(http, ref, args.sources)
        ref["source"] = {"level": level, "origin": origin, "sentences": len(sents)}
        (out / "sources" / f"ref{ref['num']}.json").write_text(json.dumps(sents, indent=0, ensure_ascii=False))
        (out / "sources" / f"ref{ref['num']}.txt").write_text(
            "\n".join(f"[s{i}] {s}" for i, s in enumerate(sents)), encoding="utf-8")
        print(f"  ref {ref['num']}: {ref['status']}, source {level}", file=sys.stderr)
        return ref

    with ThreadPoolExecutor(max_workers=4) as pool:
        refs = list(pool.map(work, refs))

    sents = body_sentences(body)
    if not numbered and any(_NUMCITE.search(s) for s in sents):
        numbered = True  # Word auto-numbered lists: the numbers are not in the text, so use list order
    unresolved = []
    citations = find_citations(sents, refs, numbered, unresolved)
    if numbered and not citations:
        citations = find_citations(sents, refs, False, unresolved)  # numbered list, author-year citations in the text
    ref_by_num = {r["num"]: r for r in refs}
    checks, missing = [], set()
    for i, nums in citations:
        for n in nums:
            ref = ref_by_num.get(n)
            if ref is None:
                missing.add(n)
                continue
            source = json.loads((out / "sources" / f"ref{n}.json").read_text())
            claim = strip_markers(sents[i])
            ranked = sorted(enumerate(bm25(claim, source)), key=lambda x: -x[1]) if source else []
            if len(source) > args.top * 2:
                ranked = ranked[:args.top]  # a short source (an abstract) is passed whole, best first
            work = ref.get("work") or {}
            if ref["status"] == "doi_mismatch" and ref.get("suggested_doi"):
                source_doi, source_title = ref.get("suggested_doi"), ref.get("suggested_title")
            else:
                source_doi, source_title = work.get("doi") or ref.get("doi"), work.get("title")
            checks.append({
                "id": f"c{i}-r{n}",
                "claim": claim,
                "claim_raw": sents[i],
                "context_before": sents[i - 1] if i else "",
                "ref": n,
                "ref_raw": ref["raw"],
                "ref_status": ref["status"],
                "ref_retracted": ref["retracted"],
                "ref_updates": work.get("updates", []),
                "ref_notes": ref["notes"],
                "source_doi": source_doi if ref["source"]["level"] != "none" else None,
                "source_title": source_title if ref["source"]["level"] != "none" else None,
                "source_level": ref["source"]["level"],
                "source_file": f"sources/ref{n}.txt",
                "candidates": [{"sid": j, "score": round(sc, 2), "text": source[j]} for j, sc in ranked],
            })
    cited = {n for _, nums in citations for n in nums}
    meta = {
        "manuscript": str(Path(args.manuscript).resolve()),
        "generated": datetime.now().isoformat(timespec="seconds"),
        "style": "numbered" if numbered else "author-year",
        "cited_but_missing": sorted(n for n in missing if n > 0),
        "unresolved_author_year": sum(1 for _, nums in citations if -1 in nums),
        "unresolved_citations": list(dict.fromkeys(unresolved)),
        "uncited_references": sorted(r["num"] for r in refs if r["num"] not in cited),
    }
    (out / "references.json").write_text(json.dumps(refs, indent=1, ensure_ascii=False))
    (out / "checks.json").write_text(json.dumps({**meta, "checks": checks}, indent=1, ensure_ascii=False))
    template = [{"id": c["id"], "verdict": "", "anchor": "", "note": ""} for c in checks]
    (out / "judgments.template.json").write_text(json.dumps(template, indent=1))
    print(f"{len(checks)} claim-reference checks written to {out / 'checks.json'}", file=sys.stderr)


# ---------------------------------------------------------------- finalize

def locate(anchor, source):
    """Find the anchor in the source. Returns (score, first_sid, text) of the best sentence or sentence pair."""
    ak = key(anchor)
    if len(ak.split()) < ANCHOR_MIN_WORDS:
        return 0, None, None
    keys = [key(s) for s in source]
    best = (0, None, None)
    for span in (1, 2):  # a sentence pair only when no single sentence holds the anchor
        for i in range(len(source) - span + 1):
            k = " ".join(keys[i:i + span])
            if len(k) < len(ak):
                continue  # a sentence shorter than the anchor would score high against part of it
            score = fuzz.partial_ratio(ak, k)
            if score > best[0]:
                best = (score, i, " ".join(source[i:i + span]))
        if best[0] >= ANCHOR_THRESHOLD:
            break
    return best


def finalize(args):
    out = Path(args.workdir)
    data = json.loads((out / "checks.json").read_text())
    refs = {r["num"]: r for r in json.loads((out / "references.json").read_text())}
    judgments = {j["id"]: j for j in json.loads((out / args.judgments).read_text())}
    rows, problems = [], []
    for c in data["checks"]:
        j = judgments.get(c["id"], {})
        verdict = (j.get("verdict") or "").strip().lower()
        anchor = (j.get("anchor") or "").strip()
        row = {"id": c["id"], "ref": c["ref"], "claim": c["claim"], "model_verdict": verdict,
               "verdict": verdict, "quote": "", "quote_sid": "", "quote_score": "", "note": j.get("note", ""),
               "source_level": c["source_level"], "ref_status": c["ref_status"],
               "source_title": c.get("source_title") or "", "source_doi": c.get("source_doi") or "",
               "code_matches": [x["text"] for x in c["candidates"][:3]],
               "code_sids": [f"s{x['sid']}" for x in c["candidates"][:3]], "in_code_top3": ""}
        if verdict not in VERDICTS:
            row["verdict"] = "not_judged"
            problems.append(f"{c['id']}: missing or unknown verdict {verdict!r}")
        elif anchor:
            source = json.loads((out / "sources" / f"ref{c['ref']}.json").read_text())
            score, sid, text = locate(anchor, source)
            row["quote_score"] = score
            if score >= ANCHOR_THRESHOLD:
                row.update(quote=text, quote_sid=f"s{sid}",
                           in_code_top3="Yes" if f"s{sid}" in row["code_sids"] else "No")
            elif verdict in NEEDS_QUOTE:
                row["verdict"] = "unverified_quote"
                problems.append(f"{c['id']}: anchor not found in source (best match {score:.0f}): {anchor!r}")
        elif verdict in NEEDS_QUOTE:
            row["verdict"] = "unverified_quote"
            problems.append(f"{c['id']}: verdict {verdict!r} given without an anchor")
        rows.append(row)

    with open(out / "report.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["id"])
        w.writeheader()
        w.writerows({k: " | ".join(v) if isinstance(v, list) else v for k, v in r.items()} for r in rows)
    with open(out / "references.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["num", "status", "retracted", "title_score", "registry_title", "doi", "suggested_doi",
                    "source", "notes", "raw"])
        for r in refs.values():
            work = r.get("work") or {}
            w.writerow([r["num"], r["status"], r["retracted"], (r.get("match") or {}).get("title_score", ""),
                        work.get("title", ""), work.get("doi") or r.get("doi") or "", r.get("suggested_doi", ""),
                        r["source"]["level"], "; ".join(r["notes"]), r["raw"]])
    (out / "report.md").write_text(render_markdown(data, refs, rows), encoding="utf-8")
    try:
        write_xlsx(out / "report.xlsx", data, refs, rows)
    except ImportError:
        print("openpyxl not installed; skipping report.xlsx", file=sys.stderr)
    (out / "report.json").write_text(json.dumps({"meta": {k: v for k, v in data.items() if k != "checks"},
                                                 "checks": rows}, indent=1, ensure_ascii=False))
    for p in problems:
        print("PROBLEM", p, file=sys.stderr)
    print(f"Report: {out / 'report.md'} ({len(problems)} problems)", file=sys.stderr)
    return 1 if problems and args.strict else 0


def _cell(s, limit=None):
    s = collapse(str(s if s is not None else ""))
    if limit and len(s) > limit:
        s = s[:limit - 1] + "…"
    return s.replace("|", "\\|")


_STATUS_TEXT = {
    "verified": "✅ verified",
    "found_without_doi": "✅ found (no DOI given)",
    "metadata_discrepancy": "⚠️ metadata differs",
    "doi_mismatch": "❌ DOI/link is a different paper",
    "doi_not_found": "❌ DOI does not exist",
    "doi_exists_no_metadata": "⚠️ DOI exists, no metadata",
    "not_found": "❌ not found",
    "lookup_failed": "⚠️ lookup failed",
}
_VERDICT_TEXT = {
    "supports": "✅ supports", "partial": "🟡 partial", "contradicts": "❌ contradicts",
    "not_found": "⚠️ not found in source", "cannot_assess": "⚪ cannot assess",
    "unverified_quote": "❓ quote not verified", "not_judged": "— not judged",
}


def render_markdown(data, refs, rows):
    ref_counts = Counter(r["status"] for r in refs.values())
    verdicts = Counter(r["verdict"] for r in rows)
    lines = [f"# Reference check: {Path(data['manuscript']).name}", "",
             f"Generated {data['generated']} · {len(refs)} references · {len(rows)} claim–reference checks "
             f"· citation style: {data['style']}", "", "## Summary", ""]
    lines += [f"- **References:** " + ", ".join(f"{_STATUS_TEXT.get(k, k)} {v}" for k, v in ref_counts.most_common())]
    retracted = [r["num"] for r in refs.values() if r["retracted"]]
    if retracted:
        lines.append(f"- **Retracted:** refs {', '.join(map(str, retracted))}")
    concern = [r["num"] for r in refs.values() if r["concern"]]
    if concern:
        lines.append(f"- **Expression of concern:** refs {', '.join(map(str, concern))}")
    lines.append("- **LLM suggested verdicts:** " + ", ".join(f"{_VERDICT_TEXT.get(k, k)} {v}" for k, v in verdicts.most_common()))
    if data["cited_but_missing"]:
        lines.append(f"- **Cited but not in reference list:** {', '.join(map(str, data['cited_but_missing']))}")
    if data.get("unresolved_citations"):
        lines.append("- **Cited in the text but not in the reference list:** "
                     + "; ".join(data["unresolved_citations"]))
    if data["uncited_references"]:
        lines.append(f"- **Never cited in the text:** refs {', '.join(map(str, data['uncited_references']))}")

    flagged = [r for r in refs.values() if r["status"] != "verified" or r["retracted"] or r["concern"]]
    for r in flagged:
        if r["status"] == "found_without_doi":
            r["notes"] = r["notes"] + [f"add DOI {r['work']['doi']}"]
    if flagged:
        lines += ["", "## References needing attention", "", "| # | Problem | Registry says | Notes |", "|---|---|---|---|"]
        for r in flagged:
            work = r.get("work") or {}
            problem = _STATUS_TEXT.get(r["status"], r["status"]) + (" · **RETRACTED**" if r["retracted"] else "")
            registry = f"{work.get('title', '')} ({work.get('doi', '')})" if work else ""
            notes = list(r["notes"]) + ([f"suggested DOI {r['suggested_doi']}"] if r.get("suggested_doi") else [])
            lines.append(f"| {r['num']} | {problem} | {_cell(registry, 160)} | {_cell('; '.join(notes))} |")

    order = ["contradicts", "unverified_quote", "not_found", "partial", "cannot_assess", "not_judged", "supports"]
    lines += ["", "## Claims", "", "| Check | Claim | Ref | LLM suggested verdict | LLM best match (verified in source) | Source |",
              "|---|---|---|---|---|---|"]
    for r in sorted(rows, key=lambda r: (order.index(r["verdict"]) if r["verdict"] in order else 99, r["ref"])):
        quote_ = f"“{r['quote']}” ({r['quote_sid']})" if r["quote"] else ""
        if r["note"]:
            quote_ += (" — " if quote_ else "") + f"*{r['note']}*"
        lines.append(f"| {r['id']} | {_cell(r['claim'], 220)} | {r['ref']} | {_VERDICT_TEXT.get(r['verdict'], r['verdict'])} "
                     f"| {_cell(quote_, 400)} | {r['source_level']} |")

    lines += ["", "## All references", "", "| # | Status | Title match | Source text | Reference |", "|---|---|---|---|---|"]
    for r in refs.values():
        score = (r.get("match") or {}).get("title_score", "")
        lines.append(f"| {r['num']} | {_STATUS_TEXT.get(r['status'], r['status'])} | {score} | "
                     f"{r['source']['level']} | {_cell(r['raw'], 200)} |")
    lines += ["", "*Verdicts here are the LLM's suggestions. The final verdict is yours: enter it in the "
              "“Your verdict” column of report.xlsx, where the code's top 3 source sentences are shown too.*"]
    lines += ["", "*Suggestions other than “supports” are worth a close look. “abstract” means only the abstract was "
              "available, so a “not found” there is weak evidence.*"]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------- Excel workbook

_VERDICT_LABEL = {
    "supports": "Supports", "partial": "Partial", "contradicts": "Contradicts", "not_found": "Not found in source",
    "cannot_assess": "Cannot assess", "unverified_quote": "Quote not verified", "not_judged": "Not judged",
}
_STATUS_LABEL = {
    "verified": "Verified", "found_without_doi": "Found (no DOI given)", "metadata_discrepancy": "Metadata differs",
    "doi_mismatch": "DOI/link is a different paper", "doi_not_found": "DOI does not exist",
    "doi_exists_no_metadata": "DOI exists, no metadata", "not_found": "Not found", "lookup_failed": "Lookup failed",
}
_GOOD, _WARN, _BAD, _GREY = "E2EFDA", "FFF2CC", "F8CBAD", "EDEDED"
_VERDICT_FILL = {"Supports": _GOOD, "Partial": _WARN, "Not found in source": _WARN, "Contradicts": _BAD,
                 "Quote not verified": _BAD, "Cannot assess": _GREY, "Not judged": _GREY}
_STATUS_FILL = {"Verified": _GOOD, "Found (no DOI given)": _GOOD, "Metadata differs": _WARN,
                "DOI exists, no metadata": _WARN, "Lookup failed": _WARN, "DOI/link is a different paper": _BAD,
                "DOI does not exist": _BAD, "Not found": _BAD}


def write_xlsx(path, data, refs, rows):
    """Results (one plain table), References, and Summary (live COUNTIF formulas)."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    review_fill = PatternFill("solid", fgColor="FFFF00")
    wrap = Alignment(wrap_text=True, vertical="top")

    def table(ws, headers, widths, records):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical="center")
        for rec in records:
            ws.append(rec)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font, cell.alignment = base, wrap
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    def color_by(ws, col, fills, last_row):
        """Row colour follows the label in `col`, so it updates if a reviewer edits it."""
        rng = f"A2:{ws.cell(1, ws.max_column).column_letter}{max(last_row, 2)}"
        for label, color in fills.items():
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col}2="{label}"'],
                                                           fill=PatternFill("solid", bgColor=color)))

    wb = Workbook()
    # Results first: one plain table, one row per claim, in three blocks:
    # what code found, what the LLM suggests, and your verdict (the one that counts)
    results = wb.active
    results.title = "Results"
    level_text = {"fulltext": "Full text", "abstract": "Abstract only", "registry": "Trial registration",
                  "none": "Nothing available"}
    records = []
    for i, r in enumerate(sorted(rows, key=lambda r: (int(r["id"].split("-")[0][1:]), r["ref"])), start=1):
        ref = refs[r["ref"]]
        check = _STATUS_LABEL.get(ref["status"], ref["status"]) + (" · RETRACTED" if ref["retracted"] else "")
        doi = ref.get("suggested_doi") or ((ref.get("work") or {}).get("doi", "")
                                           if ref["status"] == "found_without_doi" else "")
        checked = level_text.get(r["source_level"], r["source_level"])
        if r["source_title"]:
            checked += f": {r['source_title']}"
        matches = (r["code_matches"] + ["", "", ""])[:3]
        records.append([i, r["claim"], r["ref"], ref["raw"], check, doi, checked, *matches,
                        r["quote"], r["in_code_top3"], _VERDICT_LABEL.get(r["verdict"], r["verdict"]), r["note"],
                        None, None])
    table(results,
          ["#", "Claim in manuscript", "Ref #", "Reference as cited", "Reference check", "DOI to use",
           "Checked against", "Code match 1", "Code match 2", "Code match 3",
           "LLM best match (verified in source)", "LLM match in code top 3?", "LLM suggested verdict",
           "LLM comment", "Your verdict", "Your comment"],
          [5, 45, 6, 36, 20, 22, 30, 45, 45, 45, 50, 12, 16, 36, 18, 30],
          records)
    n_claims = len(rows) + 1
    last = max(n_claims, 2)
    results.freeze_panes = "C2"
    group_fill = {"code": "1F4E78", "llm": "5B3F8C", "you": "7F6000"}
    for col, group in (("A", "code"), ("B", "code"), ("C", "code"), ("D", "code"), ("E", "code"), ("F", "code"),
                       ("G", "code"), ("H", "code"), ("I", "code"), ("J", "code"),
                       ("K", "llm"), ("L", "llm"), ("M", "llm"), ("N", "llm"), ("O", "you"), ("P", "you")):
        results[f"{col}1"].fill = PatternFill("solid", fgColor=group_fill[group])
    for col in ("M", "O"):
        for label, color in _VERDICT_FILL.items():
            results.conditional_formatting.add(f"{col}2:{col}{last}", FormulaRule(
                formula=[f'${col}2="{label}"'], fill=PatternFill("solid", bgColor=color)))
    for label, color in _STATUS_FILL.items():
        results.conditional_formatting.add(f"E2:E{last}", FormulaRule(
            formula=[f'LEFT($E2,{len(label)})="{label}"'], fill=PatternFill("solid", bgColor=color)))
    results.conditional_formatting.add(f"E2:E{last}", FormulaRule(
        formula=['ISNUMBER(SEARCH("RETRACTED",$E2))'], fill=PatternFill("solid", bgColor=_BAD)))
    for row in results.iter_rows(min_row=2, min_col=13, max_col=13):
        for cell in row:
            cell.font = Font(name="Arial", size=10, italic=True, color="595959")  # a suggestion, not the verdict
    verdict_list = ",".join(v for k, v in _VERDICT_LABEL.items() if k in VERDICTS)
    decision = DataValidation(type="list", formula1=f'"{verdict_list}"', allow_blank=True)
    results.add_data_validation(decision)
    decision.add(f"O2:O{last + 200}")
    for r in range(2, last + 1):
        results[f"O{r}"].fill = PatternFill("solid", fgColor="FFF9DB")
        results[f"P{r}"].fill = PatternFill("solid", fgColor="FFF9DB")

    references = wb.create_sheet("References")
    table(references,
          ["Ref #", "Status", "Retracted", "Expression of concern", "Title match (%)", "Registry title",
           "Registry DOI", "Suggested / missing DOI", "Source text", "Notes", "Reference as cited"],
          [7, 22, 10, 12, 10, 45, 28, 28, 11, 45, 70],
          [[r["num"], _STATUS_LABEL.get(r["status"], r["status"]), "Yes" if r["retracted"] else "No",
            "Yes" if r["concern"] else "No", (r.get("match") or {}).get("title_score"),
            (r.get("work") or {}).get("title", ""), (r.get("work") or {}).get("doi", ""),
            r.get("suggested_doi") or ((r.get("work") or {}).get("doi", "") if r["status"] == "found_without_doi"
                                       else ""),
            r["source"]["level"], "; ".join(r["notes"]), r["raw"]]
           for r in refs.values()])
    n_refs = len(refs) + 1
    color_by(references, "B", _STATUS_FILL, n_refs)
    references.conditional_formatting.add(f"C2:C{max(n_refs, 2)}", FormulaRule(
        formula=['$C2="Yes"'], fill=PatternFill("solid", bgColor=_BAD), font=Font(name="Arial", bold=True)))

    # Summary: every count is a formula over the other two sheets
    summary = wb.create_sheet("Summary")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 60
    summary["A1"] = f"Reference check: {Path(data['manuscript']).name}"
    summary["A1"].font = Font(name="Arial", size=14, bold=True)
    summary["A2"] = f"Generated {data['generated']} · citation style: {data['style']}"
    summary["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    r_rng, c_rng = f"References!$B$2:$B${max(n_refs, 2)}", f"Results!$M$2:$M${max(n_claims, 2)}"
    line = 4

    def section(title):
        nonlocal line
        summary.cell(line, 1, title).font = bold
        summary.cell(line, 2, "Count").font = bold
        line += 1

    def count(label, formula, fill=None):
        nonlocal line
        summary.cell(line, 1, label).font = base
        cell = summary.cell(line, 2, formula)
        cell.font = base
        if fill:
            summary.cell(line, 1).fill = summary.cell(line, 2).fill = PatternFill("solid", fgColor=fill)
        line += 1

    section("References")
    count("Total", f"=COUNTA({r_rng})")
    for label, fill in _STATUS_FILL.items():
        count(label, f'=COUNTIF({r_rng},"{label}")', fill)
    count("Retracted", f'=COUNTIF(References!$C$2:$C${max(n_refs, 2)},"Yes")', _BAD)
    count("Expression of concern", f'=COUNTIF(References!$D$2:$D${max(n_refs, 2)},"Yes")', _WARN)
    line += 1
    y_rng = f"Results!$O$2:$O${max(n_claims, 2)}"
    section("Your verdicts")
    count("Claim checks", f"=COUNTA({c_rng})")
    count("Reviewed", f"=COUNTA({y_rng})")
    count("Not yet reviewed", f"=COUNTA({c_rng})-COUNTA({y_rng})")
    for key_, label in _VERDICT_LABEL.items():
        if key_ in VERDICTS:
            count(label, f'=COUNTIF({y_rng},"{label}")', _VERDICT_FILL[label])
    count("You agreed with the LLM", f'=SUMPRODUCT(({y_rng}={c_rng})*({y_rng}<>""))')
    line += 1
    section("LLM suggested verdicts")
    for label, fill in _VERDICT_FILL.items():
        count(label, f'=COUNTIF({c_rng},"{label}")', fill)
    line += 1
    missing = [str(n) for n in data["cited_but_missing"]] + data.get("unresolved_citations", [])
    extras = [("Cited but not in reference list", "; ".join(missing) or "none"),
              ("Never cited in the text", ", ".join(map(str, data["uncited_references"])) or "none")]
    for label, value in extras:
        summary.cell(line, 1, label).font = base
        summary.cell(line, 3, value).font = base
        line += 1
    line += 1
    notes = [
        "How to read this workbook",
        "Results: one row per claim and reference, in three blocks of columns.",
        "  Blue headers - code: reference checks and the top 3 source sentences ranked by word overlap.",
        "  Purple headers - LLM: its best matching sentence (found in the source by code, never written by "
        "the model) and a suggested verdict.",
        "  Gold headers - you: Your verdict is the final verdict. The summary counts it, not the LLM's suggestion.",
        "Source text 'abstract' means only the abstract was available; 'Not found in source' is then weak evidence.",
        "Colours: green = fine, yellow = check, red = problem, grey = could not be checked.",
    ]
    for i, text in enumerate(notes):
        cell = summary.cell(line + i, 1, text)
        cell.font = bold if i == 0 else base
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


# ---------------------------------------------------------------- CLI

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("prepare", help="parse, verify, fetch, shortlist")
    p.add_argument("manuscript", help=".docx, .pdf, .md or .txt")
    p.add_argument("--out", required=True, help="working directory for this run")
    p.add_argument("--refs", help="reference list in a separate file")
    p.add_argument("--sources", help="folder of source papers you have, named 3.pdf, ref3.pdf, ...")
    p.add_argument("--top", type=int, default=8, help="candidate passages per claim (default 8)")
    p.add_argument("--email", help="contact email for Crossref's polite pool (optional)")
    f = sub.add_parser("finalize", help="verify anchors and write the report")
    f.add_argument("workdir")
    f.add_argument("--judgments", default="judgments.json")
    f.add_argument("--strict", action="store_true", help="exit 1 if any anchor fails")
    args = ap.parse_args()
    if args.cmd == "prepare":
        prepare(args)
    else:
        sys.exit(finalize(args))


if __name__ == "__main__":
    main()
